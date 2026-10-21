"""Build an editable PowerPoint using the user's TinyTapeOut theme.

Run from any directory: python3 presentations/src/build_kev_findings.py
All results are recomputed from saved responses by analyse_kev.py first.
"""
from __future__ import annotations

import hashlib
import json
import math
import shutil
import statistics
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from analyse_kev import ROOT, MODELS, analyse, save

DEST = ROOT / 'presentations'
OUTPUT = DEST / 'KEV_RTL_findings_20261015.pptx'
TEMPLATE = ROOT / 'pwp_example/TinyTapeOut 2026 jullie feedback.pptx'
D = analyse()
save(D)
M = [x[0] for x in MODELS]
SHORT = ['0.8B', '4B', '9B', '27B']
NAMES = ['KEV-' + s for s in SHORT]
C = dict(navy='00407A', bg='F4F8FA', pale='E3EFF5', cyan='52BDEC', teal='1D8DB0',
         blue='0B5A93', muted='4A5A66', grid='C8D3DA', gray='B5BFC3', white='FFFFFF')
MC = [C['muted'], C['blue'], C['teal'], C['navy']]
prs = Presentation(TEMPLATE)
for sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)
prs.core_properties.title = 'KEV on RTL: tests, results and limitations'
prs.core_properties.subject = 'Dataset v1; all four KEV v1.0 models; 725 held-out questions'
prs.core_properties.author = 'JEV_RTL study'
prs.core_properties.keywords = 'RTL, KEV, synthesis, compilation, area, power, timing'
W, H = prs.slide_width / 914400, prs.slide_height / 914400
TEXT_BOXES = []
SLIDE_META = []


def rgb(c):
    return RGBColor.from_string(c)


def pct(value, digits=1):
    return str(Decimal(str(value * 100)).quantize(Decimal('1.' + '0' * digits), rounding=ROUND_HALF_UP)) + '%'


def metric(key, field='accuracy'):
    return [D['models'][m]['table'][key][field] for m in M]


def hits(key):
    return [round(t * D['models'][m]['table'][key]['n']) for t, m in zip(metric(key), M)]


def box(s, x, y, w, h, fill=None, line=None, radius=False):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        try: shape.adjustments[0] = 0.06
        except (IndexError, TypeError): pass
    return shape


def txt(s, text, x, y, w, h, size=18, color=None, bold=False, align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP, margin=0, record=True):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.09
        r = p.add_run(); r.text = line
        r.font.name = 'Arial'; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = rgb(color or C['navy'])
    if record:
        TEXT_BOXES.append(dict(slide=len(prs.slides), text=str(text), x=x, y=y, w=w, h=h, size=size))
    return shape


def line(s, x1, y1, x2, y2, color=None, width=1):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color or C['grid']); sh.line.width = Pt(width)
    return sh


def arrow(s, x, y, w=.34, h=.25):
    sh = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(C['teal']); sh.line.fill.background()
    return sh


def notes(s, content, source='data/questions/dataset_v1.jsonl; results/kev/*/dataset_v1__*.jsonl; harness/score.py'):
    text = (content.strip() + '\n\nEvidence and reproducibility\n' + source +
            '\nAll numerical results use the held-out test split unless explicitly labelled otherwise.' +
            '\nDataset SHA-256: ' + D['dataset_sha256'] +
            '\nVerified calculations: presentations/data/verified_metrics.json.' +
            '\nRebuild: python3 presentations/src/build_kev_findings.py.' +
            '\nVisual style: pwp_example/TinyTapeOut 2026 jullie feedback.pptx; original theme and slide masters retained.' +
            '\nAll charts, diagrams and tables on this slide are editable PowerPoint objects.')
    ns = s.notes_slide
    if ns.notes_text_frame is None:
        # The supplied reference has no body placeholder in its notes master.
        sp = OxmlElement('p:sp')
        nv = OxmlElement('p:nvSpPr')
        prop = OxmlElement('p:cNvPr'); prop.set('id', '100'); prop.set('name', 'Speaker notes')
        nv.append(prop); nv.append(OxmlElement('p:cNvSpPr'))
        nvp = OxmlElement('p:nvPr'); ph = OxmlElement('p:ph'); ph.set('type', 'body'); ph.set('idx', '13')
        nvp.append(ph); nv.append(nvp); sp.append(nv); sp.append(OxmlElement('p:spPr'))
        body = OxmlElement('p:txBody')
        for name in ['a:bodyPr', 'a:lstStyle', 'a:p']: body.append(OxmlElement(name))
        sp.append(body); ns._element.cSld.spTree.append(sp)
    ns.notes_text_frame.text = text


def slide(title, subtitle='', section='RESULTS', bg=None):
    s = prs.slides.add_slide(prs.slide_layouts[-1])
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg or C['bg'])
    txt(s, title, .89, .65, 11.55, .68, size=31, bold=True)
    if subtitle:
        txt(s, subtitle, .89, 1.38, 11.55, .48, size=16, color=C['muted'])
    txt(s, 'KEV v1.0  •  Dataset v1  •  15 October 2026  |  ' + section,
        .89, 7.04, 10.65, .22, size=10.5, color=C['muted'])
    txt(s, str(len(prs.slides)), 11.85, 7.02, .58, .25, size=11, color=C['muted'], align=PP_ALIGN.RIGHT)
    SLIDE_META.append({'number': len(prs.slides), 'title': title, 'section': section})
    return s


def card(s, x, y, w, h, title, body, color=None, body_size=18):
    if 3.5 <= h < 3.95:
        h = 3.95
    box(s, x, y, w, h, fill=C['pale'], radius=True)
    box(s, x, y, w, .045, fill=color or C['teal'])
    compact = h < 2.0
    txt(s, title, x + .25, y + (.16 if compact else .25), w - .5, .42 if compact else .72, size=20 if compact else 21, bold=True)
    txt(s, body, x + .25, y + (.72 if compact else 1.1), w - .5, h - (.82 if compact else 1.3), size=body_size)


def takeaway(s, text, y=6.25):
    box(s, .89, y, 11.55, .66, fill=C['pale'], radius=True)
    txt(s, text, 1.08, y + .08, 11.17, .50, size=15, bold=True, valign=MSO_ANCHOR.MIDDLE)


def table(s, headers, rows, x, y, widths, h, size=16, head_size=14, row_header=False):
    t = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y),
                           Inches(sum(widths)), Inches(h)).table
    for j, width in enumerate(widths):
        t.columns[j].width = Inches(width)
    t.rows[0].height = Inches(.62)
    for i in range(1, len(rows) + 1):
        t.rows[i].height = Inches((h - .62) / len(rows))
    for i, row in enumerate([headers] + rows):
        for j, value in enumerate(row):
            cell = t.cell(i, j)
            cell.margin_left = Inches(.12); cell.margin_right = Inches(.10)
            cell.margin_top = Inches(.06); cell.margin_bottom = Inches(.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(C['navy'] if i == 0 else (C['pale'] if i % 2 else C['white']))
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
                p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.03
                for r in p.runs:
                    r.font.name = 'Arial'; r.font.size = Pt(head_size if i == 0 else size)
                    r.font.bold = i == 0 or (row_header and j == 0)
                    r.font.color.rgb = rgb(C['white'] if i == 0 else C['navy'])
            for edge in ['lnL', 'lnR', 'lnT', 'lnB']:
                border = OxmlElement('a:' + edge)
                border.append(OxmlElement('a:noFill'))
                cell._tc.get_or_add_tcPr().append(border)
    return t


def panel(s, x, y, w, h, title, desc, vals, counts=None, n=None, reference=None, labels=None, label_w=1.1,
          caption='Correct answers (%) — higher is better'):
    """Four editable horizontal percentage bars, full 0-100% scale."""
    box(s, x, y, w, h, fill=C['white'], radius=True)
    txt(s, title, x + .22, y + .18, w - .44, .55, size=18, bold=True)
    txt(s, desc, x + .22, y + .67, w - .44, .38, size=12.5, color=C['muted'])
    px, pw = x + .93, w - 1.25 - label_w
    ry, step = y + 1.18, .55
    for tick in [0, 50, 100]:
        xx = px + pw * tick / 100
        line(s, xx, ry - .08, xx, ry + 3 * step + .3, color=C['grid'], width=.65)
        txt(s, str(tick), xx - .22, y + 3.22, .44, .22, size=11,
            color=C['muted'], align=PP_ALIGN.CENTER)
    if reference:
        ref, label = reference
        xx = px + pw * ref
        ln = line(s, xx, ry - .08, xx, ry + 3 * step + .3, C['muted'], 1.2)
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    for i, val in enumerate(vals):
        yy = ry + i * step
        txt(s, SHORT[i], x + .2, yy + .025, .66, .3, size=14)
        if val > 0:
            box(s, px, yy + .06, pw * val, .26, fill=MC[i])
        label = labels[i] if labels else pct(val)
        if counts is not None:
            label += f'\n{counts[i]} of {n}'
        txt(s, label, px + pw + .13, yy - .035, label_w, .47, size=12.5)
    txt(s, caption, x + .2, y + h - .57,
        w - .4, .24, size=11, color=C['muted'])
    if reference:
        txt(s, 'Dashed line: ' + reference[1], x + .2, y + h - .29,
            w - .4, .23, size=10.5, color=C['muted'])


def two_panels(title, subtitle, left, right, takeaway_text, note):
    s = slide(title, subtitle)
    panel(s, .89, 1.99, 5.65, 4.08, **left)
    panel(s, 6.79, 1.99, 5.65, 4.08, **right)
    takeaway(s, takeaway_text)
    notes(s, note)
    return s


def cover():
    s = slide('', section='OVERVIEW', bg=C['navy'])
    # Remove the ordinary header and footer for the title slide.
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    TEXT_BOXES[:] = [b for b in TEXT_BOXES if b['slide'] != 1]
    box(s, 9.35, 1.45, 2.0, 2.0, line=C['teal'], radius=True)
    txt(s, 'RTL TOOL OUTCOMES  •  KEV v1.0', .89, 1.95, 10, .35, size=16, color=C['cyan'], bold=True)
    txt(s, 'What can KEV tell us\nabout an RTL design?', .89, 2.55, 9.1, 1.62,
        size=43, color=C['white'], bold=True)
    txt(s, 'The tests, the completed Condor runs, and what the scores mean',
        .89, 4.55, 9.9, .85, size=23, color='BFDDEC')
    txt(s, 'Four model sizes  •  151 retained designs  •  725 held-out questions\n15 October 2026',
        .89, 6.12, 11.4, .75, size=16, color='BFDDEC')
    SLIDE_META[0]['title'] = 'What can KEV tell us about an RTL design?'
    notes(s, 'This presentation covers the completed KEV-0.8B, KEV-4B, KEV-9B and KEV-27B v1.0 runs. RTL means register-transfer level: code describing digital hardware. The label B denotes approximately billions of model parameters. There are 151 retained designs across development and test; the held-out test has 123 individual designs plus pair requests. The number 725 counts questions, not independent designs. All findings are specific to this dataset and tool flow. No Jev, base Qwen or commercial LLM evaluation is included.')


cover()

# 2: executive findings.
s = slide('What the completed runs show', 'Different tasks reveal different strengths. A single pooled score hides important failures.', 'OVERVIEW')
card(s, .89, 2.0, 3.7, 3.7, 'Better with 27B', 'Compilation errors:\n8 of 9 detected\n\nArea levels:\n61 of 92 exactly right', C['navy'], 19)
card(s, 4.81, 2.0, 3.7, 3.7, 'Synthesis is weak', 'Synthesis failures:\n4 of 24 detected\n\nFailure reasons:\n15 of 16 right when\nfailure is supplied', C['teal'], 19)
card(s, 8.74, 2.0, 3.7, 3.7, 'No single winner', '9B leads area ranking.\n27B leads power ranking.\n\nAbsolute power and\nclock-frequency levels\nremain unreliable.', C['blue'], 18)
takeaway(s, 'All four runs completed. Every model answered every test question; missing data does not explain the errors.')
notes(s, 'The compilation-error result is KEV-27B on the 9 compiler-rejected test designs. The area result is 61/92 exact levels. Synthesis failure detection is 4/24. The 15/16 failure-reason result is conditional on a prompt explicitly stating that synthesis failed. The strongest observed pair-area accuracy is KEV-9B at 40/46 views; the strongest observed pair-power accuracy is KEV-27B at 31/44 views. These are different models and tasks, and not a single combined design or statistically established winner.')

# 3: tool definitions.
s = slide('Compilation and synthesis answer different questions', 'RTL = register-transfer level: source code that describes digital hardware.', 'TESTS')
card(s, .89, 2.0, 3.7, 3.7, 'Does it compile?', 'Can the simulator compiler\naccept the source code?\n\nReference: QuestaSim\nvlog -sv.\n\nThis does not test whether\nthe circuit is correct.', body_size=18)
card(s, 4.81, 2.0, 3.7, 3.7, 'Does it synthesize?', 'Can Design Compiler build\na complete gate-level\nnetlist from the code?\n\nNo unresolved modules.\nAll registers map to cells.\nNo synthesis errors.', body_size=18)
card(s, 8.74, 2.0, 3.7, 3.7, 'What does it cost?', 'Estimate cell area, power,\nand timing from the RTL.\n\nThe reference comes from\nsynthesis and power tools.\n\nNo fabricated-chip or\npost-route measurement.', body_size=18)
takeaway(s, 'In the test set, 15 designs compile successfully but still fail synthesis.')
notes(s, 'Compilation uses Siemens QuestaSim 2022.4 vlog -sv, a compile-only check. Synthesis uses Synopsys Design Compiler U-2022.12-SP2 with analyze, elaborate, link and compile_ultra. A synthesis pass requires no Error lines, no LINK-5/UID-341 unresolved references, no TRANS-4 unmapped generic registers, and an output netlist. Compiler acceptance is not functional correctness, complete simulator elaboration, or a guarantee of synthesizability. There are 99 test designs with both successful compilation and synthesis, 15 that compile but fail synthesis, and 9 that fail both.', 'notes/synth_yesno.md; notes/DATASET.md; harness/build.py; data/questions/dataset_v1.jsonl')

# 4: coverage.
s = slide('The dataset and the held-out test set', 'One model request can contain several questions. Questions are not independent design samples.', 'TESTS')
for x, value, label in [( .89, '151', 'retained designs'), (4.81, '348', 'requests across both splits'), (8.74, '875', 'questions across both splits')]:
    box(s, x, 2.0, 3.7, 1.18, C['pale'], radius=True)
    txt(s, value, x + .2, 2.13, 1.15, .55, size=32, bold=True)
    txt(s, label, x + 1.4, 2.2, 2.08, .65, size=16)
table(s, ['Held-out request group', 'Requests', 'Questions', 'What is covered'], [
    ['Basic checks', '123', '246', 'Compilation + synthesis'],
    ['Failure reason', '16', '16', 'One of ten reason options'],
    ['Performance', '92', '361', 'Area, power and two timing levels'],
    ['Pair comparisons', '50', '102', '25 distinct pairs, both orders'],
    ['TOTAL TEST', '281', '725', '123 individual test designs'],
], .89, 3.5, [3.25, 1.2, 1.25, 5.85], 2.5, size=16)
takeaway(s, 'Development: 28 designs, 67 requests, 150 questions. Related base designs stay in the same split.')
notes(s, 'The plan ran 159 designs; 8 are excluded under the dataset rules, leaving 151 basic design records. The test split contains 123 basic designs and the development split 28. All four models answered all 348 requests, but reported accuracy uses test only: 281 requests and 725 questions. Performance coverage is 92 area, 92 slack, 92 reachable-frequency, and 85 power questions. The 361 performance questions therefore do not mean 361 distinct designs. There are 25 distinct test pairs, each sent in both orders, containing 46 area, 44 power and 12 frequency questions. Pair metrics overlap in which pairs they cover. The analysis asserts no shared base_design_id between development and test.', 'notes/DATASET.md; data/questions/dataset_v1.card.md; data/questions/dataset_v1.jsonl; presentations/src/analyse_kev.py')

# 5: model inputs.
s = slide('What the model receives for each test', 'Tool logs and correct answers are used for scoring; they are not sent to the model.', 'TESTS')
card(s, .89, 2.0, 5.65, 1.9, 'Basic checks', 'RTL source + top module + tool description.\nQuestions: does it compile, and does it synthesize?', body_size=17)
card(s, 6.79, 2.0, 5.65, 1.9, 'Failure reason', 'The same source, plus a prompt saying synthesis failed.\nChoose one reason from ten descriptions.', body_size=17)
card(s, .89, 4.1, 5.65, 1.9, 'Performance', 'RTL + flow + clock target; power also gets its testbench.\nChoose a level for each available metric.', body_size=17)
card(s, 6.79, 4.1, 5.65, 1.9, 'Pair comparison', 'Both RTL variants + the same flow and clock.\nPower comparisons also share the same testbench.', body_size=17)
takeaway(s, 'The failure-reason task supplies a key fact that the synthesis yes/no task must discover.')
notes(s, 'The Condor client sends only state, API model and questions to /v1/systemone. The correct_answer and evidence fields are retained locally in the input file and output records for scoring, not sent in the request payload. The four request groups are deliberately separate. In particular the reason instruction literally states that the tool does not synthesize the module, so its score cannot be read as autonomous failure detection. All models used the same input file SHA-256 and v1.0-tagged checkpoints. Maximum actual reported input length was 7,706 tokens for every model.', 'condor/kev_client.py; harness/build.py: V1_QUESTIONS and build_dataset_v1; data/questions/dataset_v1.jsonl')

# 6: scoring.
s = slide('How to read the scores', 'Every result needs a denominator and a clear rule for what counts as correct.', 'TESTS')
card(s, .89, 2.0, 3.7, 3.72, 'Yes or no', 'A yes probability of\n0.5 or more means yes.\n\nBalanced accuracy gives\n“yes” and “no” designs\nequal weight.\n\nAlways one answer: 50%.', body_size=18)
card(s, 4.81, 2.0, 3.7, 3.72, 'Five levels', 'Exact: the chosen level\nmatches the tool label.\n\nWithin one: a difference\nof at most one level.\n\nThis does not mean a fixed\nphysical error tolerance.', body_size=18)
card(s, 8.74, 2.0, 3.7, 3.72, 'Choose A or B', 'Each pair is sent twice,\nwith its order swapped.\n\nWe report per-view accuracy\nand check both orders.\n\n50% is an always-A or\nalways-B reference here.', body_size=18)
takeaway(s, 'A “constant-answer reference” below means always choosing one stated answer, calculated from test labels.')
notes(s, 'Binary predictions use threshold 0.5 exactly as the saved scorer. Class-specific rates use the actual tool outcome as the grouping variable. Binary balanced accuracy is the arithmetic mean of success recognition and failure detection. It equals the accuracy on a test set with equal numbers of “yes” and “no” designs, so always “yes” and always “no” both score 50% whatever the mix of the test set. The yes/no results slide shows this one score per model, with the answer mix of the model beside each bar. For reason labels the scorer averages per-class recall over the eight true classes present. Score probabilities use KEV keys 0 to 4, translated to displayed levels 1 to 5 before scoring. A constant-answer reference is a descriptive frequency of the specified label in this test set; it was not a separate run, not learned on development, and not chosen as a deployable predictor. Pair labels are exactly balanced A/B because both orders are included; always choosing a letter therefore gets 50%. There is no grand overall score across unrelated tasks.')

# 7: both yes/no checks, one score per model, with the answer mix as the constant-answer check.
YES = {key: [D['models'][m]['binary'][key]['yes_answers'] for m in M] for key in ['compiles', 'synthesizes']}
SYN_P = {m: [p for _, p in D['models'][m]['noul']['synthesizes']['pass']] for m in M}
SYN_F = {m: [p for _, p in D['models'][m]['noul']['synthesizes']['fail']] for m in M}
COMP_08 = [p for g in ['pass', 'fail'] for _, p in D['models']['kev-0.8b']['noul']['compiles'][g]]


def yes_no_panel(key, title):
    b = D['models'][M[0]]['binary'][key]
    n = b['n_pass'] + b['n_fail']
    vals = metric(key, 'balanced_accuracy')
    return dict(title=title, desc=f'Tool answers for the {n} designs: {b["n_pass"]} yes · {b["n_fail"]} no', vals=vals,
                labels=[f'{pct(v)}\n{y} yes · {n - y} no' for v, y in zip(vals, YES[key])], label_w=1.35,
                reference=(.5, 'always “yes” or always “no” = 50%'), caption='Balanced accuracy (%) — higher is better')


# The title and the takeaway are checked here: one score is far above 50%, and every other
# answer set gives one answer to at least 118 of the 123 designs.
SAME = {(m, key): max(y, 123 - y) for key in YES for m, y in zip(M, YES[key])}
assert all(D['models'][m]['table'][key]['balanced_accuracy'] <= .62 for m, key in SAME if (m, key) != ('kev-27b', 'compiles'))
assert D['models']['kev-27b']['table']['compiles']['balanced_accuracy'] > .9
assert min(v for k, v in SAME.items() if k != ('kev-27b', 'compiles')) >= 118 > SAME[('kev-27b', 'compiles')]
two_panels('Only KEV-27B’s compilation score is far above 50%', 'One score per model: “yes” and “no” designs get equal weight. A constant answer scores 50%.',
    yes_no_panel('synthesizes', 'Does it synthesize?'), yes_no_panel('compiles', 'Does it compile?'),
    'In the other seven cases, at least 118 of the 123 answers are the same.',
    'One test set of 123 held-out designs answers both questions; all designs were sent together. Each bar is one score per model and question: balanced accuracy, the mean of the share of “yes” designs answered “yes” and the share of “no” designs answered “no”. '
    'It equals the accuracy on a test set with equal numbers of “yes” and “no” designs, so any constant answer scores exactly 50%. Plain accuracy over all 123 designs rewards a constant “yes”: always “yes” scores 80.5% for synthesis (99 of 123) and 92.7% for compilation (114 of 123). '
    'Beside each bar is the answer mix of the model (a yes probability of 0.5 or more counts as “yes”); compare it with the answer mix of the tool above the chart. '
    'Synthesis: “yes” designs answered “yes” 99, 98, 97, 98 of 99; “no” designs answered “no” 0, 0, 2, 4 of 24 (0.8B, 4B, 9B, 27B). Compilation: 4, 114, 114, 114 of 114 and 9, 0, 2, 8 of 9. '
    'With only 9 compile errors, one design changes the balanced score by 5.6 percentage points. '
    f'KEV-0.8B answers “no” to compilation for {123 - YES["compiles"][0]} of the 123 designs, with yes probabilities between {min(COMP_08):.2f} and {max(COMP_08):.2f} for every design: its 9 of 9 on the compile errors is this near-constant “no”, not a detection. '
    'The question is yes/no only; it does not ask which error occurred (that is the failure-reason test, where 0.8B is right for 5 of 16). '
    f'The same KEV-0.8B answers “yes” to synthesis for all {YES["synthesizes"][0]} designs, so for {D["models"]["kev-0.8b"]["no_compile_but_synthesizes"]} designs it answers both “does not compile” and “synthesizes”; in this test set, every design that vlog rejects also fails synthesis. '
    'The appendix slide “Most failing designs get the same answer as successes” shows every yes probability.')

# 8: explicit confusion matrix.
s = slide('Why 82.9% hid poor synthesis failure detection', 'KEV-27B • 123 held-out designs • a yes/no threshold of 0.5', 'RESULTS')
table(s, ['Actual tool result', 'Model says YES', 'Model says NO'], [
    ['Synthesizes — 99 designs', '98 correct', '1 wrong'],
    ['Fails synthesis — 24 designs', '20 wrong', '4 correct'],
], .89, 2.05, [3.1, 2.55, 2.55], 1.9, size=21, head_size=17)
card(s, .89, 4.2, 5.65, 1.78, '82.9% overall accuracy', '102 correct of 123 designs.\nEach design has equal weight; the 99 successes dominate.', body_size=17)
card(s, 6.79, 4.2, 5.65, 1.78, '57.8% balanced accuracy', 'Average of 99.0% success recognition\nand 16.7% failure detection. Each group gets half the weight.', body_size=17)
txt(s, '20 failures\nmissed', 9.55, 2.1, 2.6, 1.5, size=31, bold=True)
takeaway(s, 'Both percentages are correct. They answer different questions about an imbalanced test set.')
notes(s, 'KEV-27B confusion counts: true positive 98, false negative 1, false positive 20, true negative 4. Overall accuracy = (98+4)/123 = 0.8292682927. Success recognition = 98/99 = 0.9898989899. Failure detection = 4/24 = 0.1666666667. Balanced accuracy = (98/99+4/24)/2 = 0.5782828283. The always-yes descriptive reference is 99/123 = 80.5% overall and 50.0% balanced. The matrix contains actual counts, not probabilities. Never interpret overall 82.9% as 82.9% of failures detected.')

# 9: synthesis-only errors.
left = [D['models'][m]['synthesis_failure_subsets']['compiler_accepts'] for m in M]
right = [D['models'][m]['synthesis_failure_subsets']['compiler_rejects'] for m in M]
two_panels('Which synthesis failures do the models detect?', 'Split the 24 failures again: source accepted by the simulator compiler versus rejected source.',
    dict(title='Compiles, but fails synthesis', desc='15 designs • synthesis-specific failure cases',
         vals=[x['correct']/x['n'] for x in left], counts=[x['correct'] for x in left], n=15),
    dict(title='Compilation also fails', desc='9 designs • both tools reject the design',
         vals=[x['correct']/x['n'] for x in right], counts=[x['correct'] for x in right], n=9),
    'KEV-27B misses every one of the 15 synthesis failures whose source compiles.',
    'This is a derived subgroup analysis using the actual compilation and synthesis labels of the same basic requests. All bars are synthesis yes/no decisions; the right panel is not the separate compilation answer. Compiler-accepted failure detections: 0,0,1,0 of 15. Compiler-rejected failure detections: 0,0,1,4 of 9. Small subgroups do not establish general performance, but clearly localise the observed weakness.')

# Helpers for the appendix slide with every yes probability.
OK_COLOR, FAIL_COLOR = C['teal'], 'E97132'  # validated pair on white: CVD dE 17.3, contrast 3.83 and 3.06


def dot(s, cx, cy, d, color):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color)
    sh.line.color.rgb = rgb(C['white']); sh.line.width = Pt(.6)
    return sh


def legend(s, x, y, w, items, size=12.5):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(.3))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    for i, (color, label) in enumerate(items):
        for text, col in [(('      ' if i else '') + '●', color), (' ' + label, C['muted'])]:
            r = p.add_run(); r.text = text
            r.font.name = 'Arial'; r.font.size = Pt(size); r.font.color.rgb = rgb(col)
    TEXT_BOXES.append(dict(slide=len(prs.slides), text='   '.join(l for _, l in items), x=x, y=y, w=w, h=.3, size=size))


def jitter(qid):
    """Fixed vertical offset in [-1, 1] per question ID, so every build draws the same dots."""
    return int.from_bytes(hashlib.sha256(qid.encode()).digest()[:4], 'big') / 2**31 - 1


def strip_panel(s, x, y, w, h, title, key, labels):
    box(s, x, y, w, h, fill=C['white'], radius=True)
    txt(s, title, x + .22, y + .18, w - .44, .55, size=18, bold=True)
    n = D['models'][M[0]]['binary'][key]
    legend(s, x + .22, y + .67, w - .44, [(OK_COLOR, f'{labels[0]}: {n["n_pass"]}'), (FAIL_COLOR, f'{labels[1]}: {n["n_fail"]}')])
    px, pw = x + .93, w - 2.35
    ry, step = y + 1.18, .55
    for tick in [0, .5, 1]:
        xx = px + pw * tick
        ln = line(s, xx, ry - .08, xx, ry + 3 * step + .4, color=C['muted'] if tick == .5 else C['grid'], width=1.2 if tick == .5 else .65)
        if tick == .5:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        txt(s, f'{tick:g}', xx - .22, y + 3.22, .44, .22, size=11, color=C['muted'], align=PP_ALIGN.CENTER)
    for i, m in enumerate(M):
        yy = ry + i * step
        txt(s, SHORT[i], x + .2, yy + .025, .66, .3, size=14)
        for lane, group, color in [(yy + .08, 'pass', OK_COLOR), (yy + .30, 'fail', FAIL_COLOR)]:
            for qid, p in D['models'][m]['noul'][key][group]:
                dot(s, px + pw * p, lane + .055 * jitter(qid), .08, color)
        txt(s, f'{D["models"][m]["binary"][key]["yes_answers"]} of 123\nare “yes”', px + pw + .13, yy - .035, 1.1, .47, size=12.5)
    txt(s, 'Yes probability of each design • dashed line: 0.5 threshold', x + .2, y + h - .57, w - .4, .24, size=11, color=C['muted'])


# 10: failure-reason test definition: the ten options and how often each one is correct.
RO, RC = D['reason_options'], D['reason_counts_by_split']
NT, ND = sum(RC['test'].values()), sum(RC['dev'].values())
NR = D['synthesis_failures_without_reason']
s = slide('Failure reason: choose one of ten options', f'The prompt says that synthesis failed. {NT + ND} reason questions: {NT} test, {ND} development (dev).', 'TESTS')
t = table(s, ['Option', 'Text of the option, as the model reads it', 'Test', 'Dev'],
          [[k, v, RC['test'].get(k, 0), RC['dev'].get(k, 0)] for k, v in RO.items()],
          .89, 1.95, [2.3, 7.65, .7, .9], 4.0, size=10.5, head_size=13, row_header=True)
t.rows[0].height = Inches(.4)
for i, v in enumerate(RO.values(), 1):
    t.rows[i].height = Inches(.25 if len(v) <= 100 else .40)
    for p in t.cell(i, 1).text_frame.paragraphs:
        p.alignment = PP_ALIGN.LEFT
for p in t.cell(0, 1).text_frame.paragraphs:
    p.alignment = PP_ALIGN.LEFT
assert NR['splits'] == {'test': NR['n']} and NR['classes'] == {'illegal_code': NR['n']} == {'illegal_code': NR['compile_fails']}
takeaway(s, f'{D["reason_injected_edits"]} of the {NT + ND} failures are injected edits. {NR["n"]} more test designs with a compile error get no reason question.')
notes(s, f'This slide lists the ten options of the failure-reason question exactly as the model reads them (option name, then its text). Every one of the {NT + ND} reason questions in dataset v1 offers the same ten options in the same order. '
    f'The columns count the reason questions whose correct answer is that option: {NT} in the test split and {ND} in the development split. In the test split, {len(RC["test"])} of the ten options occur; event_list and unbounded_loop occur only in the development split, once each. '
    'The correct answer comes from the Design Compiler message IDs (and the vlog result to separate illegal_code from unsupported_construct), with the rules in notes/synth_reason.md; a design is kept only if all its messages belong to one class. '
    f'{D["reason_injected_edits"]} of the {NT + ND} failing designs are injected one-construct edits of a working design (mutants and the rejected twins of the trap pairs). The other {NT + ND - D["reason_injected_edits"]} are natural failures: one MetRex design and two LLM-written candidates. '
    f'{NR["n"]} further test designs fail synthesis with illegal_code; they are the injected compile errors (one removed semicolon or one undeclared name) and get only the yes/no questions, by the dataset rules.',
    'data/questions/dataset_v1.jsonl (records with request = reason); notes/synth_reason.md; notes/DATASET.md; presentations/data/verified_metrics.json')

# 11: reason.
s = slide('Failure reasons: 27B gets 15 of 16 right', 'The prompt explicitly states that synthesis failed, then offers ten possible reasons.', 'RESULTS')
panel(s, .89, 2.0, 5.65, 4.08, 'Choose the correct failure reason', '16 known-failure designs • one choice per design', metric('reason'), hits('reason'), 16)
card(s, 6.79, 2.0, 5.65, 2.05, 'What is measured', 'Can the model identify the reason\nwhen failure is already known?', body_size=21)
card(s, 6.79, 4.25, 5.65, 1.82, 'What the 27B error was', 'A multiple-driver error was labelled\nas an asynchronous-reset error.', body_size=19)
takeaway(s, 'Only eight of the ten reason classes occur in these 16 test designs. Coverage is limited.')
notes(s, 'Correct reason counts are 5,9,11,15 of 16. Class-balanced reason accuracies are 17.9%,37.5%,50.0%,87.5%, respectively. The reason classes present are multiple_drivers(1), illegal_code(1), unsupported_construct(1), timing_control(3), simulation_statement(5), async_reset(1), mixed_assignments(2), instance_mismatch(2). There are no true event_list or unbounded_loop examples in the held-out reason subset. KEV-27B gets only multiple_drivers wrong, predicting async_reset. The instruction supplies the known failure and therefore this accuracy is conditional, not an end-to-end detection score.')

# 12: linked result.
s = slide('Diagnosing a known failure is not detecting that failure', 'Match the two saved responses for the same 16 reason-test designs.', 'RESULTS')
table(s, ['Model', 'Says synthesis failed', 'Correct reason when told it failed', 'Both answers correct'],
      [[NAMES[i], f"{D['models'][m]['reason_linked']['detected']} of 16", f"{D['models'][m]['reason_linked']['reason_correct']} of 16", f"{D['models'][m]['reason_linked']['both']} of 16"] for i,m in enumerate(M)],
      .89, 2.1, [1.5, 3.0, 4.05, 3.0], 2.75, size=21, head_size=16)
box(s, .89, 5.15, 11.55, .92, fill=C['pale'], radius=True)
txt(s, '27B: 15 correct diagnoses, but zero failure flags on these same designs.', 1.15, 5.37, 11.0, .55, size=23, bold=True)
takeaway(s, 'This is a linked analysis of existing answers, not a separately run detection-and-diagnosis system.')
notes(s, 'For each test reason request, replace :reason: with :basic: in its question ID and inspect the corresponding synthesis probability. The reason subset has 16 known failures, while the full binary synthesis test has 24 failures; 8 noncompile designs do not receive a reason request. Failure flags within the 16-design reason subset are 0,0,1,0. Correct diagnoses when the prompt supplies failure are 5,9,11,15. Both correct counts are 0,0,1,0. In particular KEV-27B flags none of these 16, although its independent reason response is right for 15. These are retrospective conjunctions of saved answers, not measured performance of a deployed routed pipeline.')

# 13: area / power test definition.
s = slide('Area and power: choose one of five ranges', 'The model estimates a level from source code and flow information, not an exact physical value.', 'TESTS')
table(s, ['Metric', 'Level 1', 'Level 2', 'Level 3', 'Level 4', 'Level 5'], [
    ['Cell area\n(gate equivalents)', '<100', '100 to\n<1,000', '1,000 to\n<10,000', '10,000 to\n<100,000', '≥100,000'],
    ['Total power\n(microwatts)', '<10', '10 to\n<100', '100 to\n<1,000', '1,000 to\n<10,000', '≥10,000'],
], .89, 2.1, [2.35, 1.84, 1.84, 1.84, 1.84, 1.84], 2.3, size=18, head_size=16)
card(s, .89, 4.7, 5.65, 1.35, 'Area reference', '1 gate equivalent = one two-input NAND gate\n= 0.196608 square micrometres in this library.', body_size=17)
card(s, 6.79, 4.7, 5.65, 1.35, 'Power reference', 'Dynamic + leakage, for the supplied stimulus.\n1,000 microwatts = 1 milliwatt.', body_size=17)
takeaway(s, 'Exact level is strict; “within one level” can span a large physical range because these bins are logarithmic.')
notes(s, 'The reference synthesis uses the supplied RTL_flow_Nangate15, NanGate 15nm typical corner at 0.8 V and 25 C, Design Compiler compile_ultra -gate_clock, a 500 MHz target, input/output delay of 10% of period, and 50 ps setup clock uncertainty. Area is the synthesis total cell area divided by 0.196608 square micrometres per NAND2_X1 gate. Power comes from the same design run and PrimeTime total dynamic plus leakage power using its supplied power testbench. The model receives that testbench. Power is only included when the run is successful and switching annotation is at least 95%. In the held-out test 85 of 92 performance requests have a power label; the other seven are excluded, not model errors. Levels use lower-inclusive upper-exclusive ranges, except the first and last open-ended bins.', 'harness/build.py: AREA_EDGES_GE, POWER_EDGES_UW and build_dataset_v1; notes/DATASET.md; notes/README.md')

# 14: timing test definition.
s = slide('Timing: margin at 500 MHz, plus a derived frequency', 'The clock target is 500 MHz: one cycle lasts 2 nanoseconds (ns).', 'TESTS')
table(s, ['Metric', 'Level 1', 'Level 2', 'Level 3', 'Level 4', 'Level 5'], [
    ['Worst slack\n(timing margin)', 'Negative:\nmisses target', '0 to\n<0.5 ns', '0.5 to\n<1.0 ns', '1.0 to\n<1.5 ns', '≥1.5 ns'],
    ['Reachable clock\nfrequency', '<1 GHz', '1 to\n<2 GHz', '2 to\n<3 GHz', '3 to\n<4 GHz', '≥4 GHz'],
], .89, 2.05, [2.35, 1.84, 1.84, 1.84, 1.84, 1.84], 2.25, size=18, head_size=16)
txt(s, 'Reachable frequency (GHz) = 1 ÷ [2 ns − worst slack (ns)]', 1.05, 4.66, 11.1, .55, size=25, bold=True)
txt(s, 'Example: 1.438 ns of slack gives a 0.562 ns limiting delay, or about 1.78 GHz.\nThis is derived from the netlist synthesized at 500 MHz; it is not a new optimized run.',
    1.05, 5.38, 11.1, .75, size=18)
takeaway(s, 'No negative-slack designs occur in this test split. Timing-violation detection is therefore not evaluated.')
notes(s, 'Worst slack is the smallest timing margin reported after synthesis at the fixed 500 MHz target. Positive slack means the netlist meets that target. Reachable frequency is the harness formula 1000/(period_ns-slack_ns) MHz, converted here to GHz, from that same netlist. It is not measured silicon frequency, post-route timing or the result of resynthesizing at the inferred faster target. The example mac_int8 has 1437.97 ps of slack, converted to 1.43797 ns: 1/(2-1.43797)=1.779 GHz. A unit conversion bug in the supplied summary.sh is corrected by the harness. Held-out slack truths contain levels 2,3,4,5 only; reachable frequency truths contain levels 1,2,3,4 only. Thus no negative-slack or ≥4 GHz true examples occur.', 'notes/DATASET.md; harness/build.py: SLACK_EDGES_NS, REACH_EDGES_MHZ and build_dataset_v1; flow/RTL_flow_Nangate15/summary.sh')

# 15: area.
two_panels('Area levels improve substantially with 9B and 27B', '92 held-out designs • same flow and target • exact and within-one scores shown separately.',
    dict(title='Exact area level', desc='All five true area levels occur', vals=metric('area'), counts=hits('area'), n=92,
         reference=(33/92, 'always level 1 = 35.9%')),
    dict(title='Within one area level', desc='Chosen level differs by at most one', vals=metric('area','within_one'),
         counts=[round(v*92) for v in metric('area','within_one')], n=92),
    '27B: 61 of 92 exact levels; all 92 estimates are within one level.',
    'Exact area hits: 33,33,53,61 of 92. Within-one hits: 61,65,83,92. The dashed constant-answer reference is always area level 1, correct for 33/92 test examples. It is descriptive and computed from test labels, not an independently run or development-selected model. Even a within-one prediction can imply a large area difference because the bins have decade-wide boundaries. No physical relative error or post-layout area is reported.')

# 16: power.
two_panels('Power levels: the largest model performs worst', '85 held-out designs with usable power measurements and a supplied stimulus testbench.',
    dict(title='Exact power level', desc='Levels 1 to 4 occur; level 5 is absent', vals=metric('power'), counts=hits('power'), n=85,
         reference=(37/85, 'always level 2 = 43.5%')),
    dict(title='Within one power level', desc='Chosen level differs by at most one', vals=metric('power','within_one'),
         counts=[round(v*85) for v in metric('power','within_one')], n=85),
    '27B predicts only levels 3–5, although 60 of the 85 true labels are levels 1–2.',
    'Exact power hits: 23,39,34,9 of 85. Within-one hits: 70,78,79,55. True level counts 1–5: 23,37,18,7,0. KEV-27B chosen-level counts: 0,0,54,30,1. It therefore systematically chooses high levels on this test distribution. This is an observed tendency, not evidence of an identified root cause. Power depends on both the synthesized design and stimulus; it is not a general intrinsic cost of the RTL. The constant-answer level-2 reference is computed from test labels.')

# 17: timing.
two_panels('Timing levels remain unreliable for every model', '92 held-out designs • these are level estimates from RTL, not timing-tool replacements.',
    dict(title='Exact clock-frequency level', desc='Frequency derived from the 500 MHz netlist', vals=metric('reach'), counts=hits('reach'), n=92,
         reference=(45/92, 'always level 4 = 48.9%')),
    dict(title='Exact timing-margin level', desc='Worst slack at the 500 MHz target', vals=metric('slack'), counts=hits('slack'), n=92,
         reference=(63/92, 'always level 5 = 68.5%')),
    '27B’s better slack result still falls below the stated constant-answer reference.',
    'Reachable frequency exact hits are 13,11,13,13 of 92. Slack exact hits are 8,4,3,30 of 92. True frequency level 4 is most common with 45/92 and true slack level 5 with 63/92. These are explicitly specified constant-answer references calculated from the test labels, not learned predictive baselines. No negative-slack cases are present, so the benchmark cannot evaluate missed timing violations. Level accuracy is not a timing estimation error in picoseconds or MHz.')

# 18: 27B distribution.
def distribution(s, x, y, w, title, truth, pred, n):
    box(s, x, y, w, 4.05, C['white'], radius=True)
    txt(s, title, x+.22, y+.18, w-.44, .38, size=19, bold=True)
    txt(s, f'{n} designs • counts at each level', x+.22, y+.67, w-.44, .3, size=13, color=C['muted'])
    px, py, pw, ph = x+.63, y+1.16, w-.98, 2.08
    for tick in [0,25,50,75,100]:
        yy=py+ph*(1-tick/100)
        line(s,px,yy,px+pw,yy,C['grid'],.65)
        txt(s,str(tick),x+.07,yy-.11,.45,.23,size=11,color=C['muted'],align=PP_ALIGN.RIGHT)
    slot=pw/5
    for i,(a,b) in enumerate(zip(truth,pred)):
        xx=px+i*slot+.08
        for j,(v,col) in enumerate([(a,C['cyan']),(b,C['navy'])]):
            if v: box(s,xx+j*.29,py+ph*(1-v/100),.22,ph*v/100,col)
            txt(s,str(v),xx+j*.29-.03,py+ph*(1-v/100)-.27,.29,.22,size=11,align=PP_ALIGN.CENTER)
        txt(s,str(i+1),xx+.05,py+ph+.1,.4,.25,size=13,align=PP_ALIGN.CENTER)
    txt(s,'Level (1–5)',x+.22,y+3.6,1.65,.28,size=12,color=C['muted'])
    box(s,x+2.15,y+3.6,.16,.16,C['cyan']); txt(s,'Tool',x+2.38,y+3.55,.6,.27,size=12)
    box(s,x+3.15,y+3.6,.16,.16,C['navy']); txt(s,'27B',x+3.38,y+3.55,.65,.27,size=12)
s=slide('27B’s predicted levels reveal the mismatch', 'The chart heights are design counts, not accuracy. Every level is shown, including zero counts.', 'RESULTS')
distribution(s,.89,2.0,5.65,'Clock-frequency levels',[13,16,18,45,0],[90,0,0,0,2],92)
distribution(s,6.79,2.0,5.65,'Power levels',[23,37,18,7,0],[0,0,54,30,1],85)
takeaway(s,'27B usually chooses the lowest frequency level and high power levels; the tool labels have a different distribution.')
notes(s,'These are the complete true-label and chosen-label distributions for KEV-27B test performance responses. Frequency truths 13,16,18,45,0 and predictions 90,0,0,0,2. Power truths 23,37,18,7,0 and predictions 0,0,54,30,1. Each chart uses the same 0–100-design vertical scale; the two denominators are different and labelled. There is no omission or zoom. Counts at zero are printed as zero rather than represented by a nonzero bar. These distributions establish bias in chosen levels, not its cause.')

# 19: pair protocol.
s=slide('Pair tests ask which variant is better', 'Two candidates for the same problem, with the same flow and power stimulus.', 'TESTS')
card(s,.89,2.0,5.65,2.65,'Three comparisons','Smaller cell area\nLower total power\nHigher derived clock frequency',body_size=23)
card(s,6.79,2.0,5.65,2.65,'Eligibility and coverage','Metric differs by at least 10%.\nArea: 23 distinct pairs, 46 views.\nPower: 22 pairs, 44 views.\nFrequency: 6 pairs, 12 views.',body_size=19)
txt(s,'First request:  [design X = A]  [design Y = B]',1.1,5.02,10.95,.4,size=21,bold=True)
txt(s,'Swap request:  [design Y = A]  [design X = B]',1.1,5.65,10.95,.4,size=21,bold=True)
takeaway(s,'A model tracking the better design should switch letters when the design order is swapped.')
notes(s,'Pair inputs are old REvolution candidates of the same problem. Both runs must have usable power and at least 95% annotation, the same power testbench, and positive values for compared metrics. A question is retained for a metric if max/min is at least 1.1, so there are no ties. Twenty-five distinct pairs occur in the held-out split and 31 across both splits. The metric subsets overlap; do not sum 23+22+6 as distinct pairs. Each retained pair is sent as AB and BA. A/B labels are therefore balanced. No energy-per-operation pair question is implemented in dataset v1, despite its presence in an older planning note.', 'harness/build.py: V1_PAIR_QUESTIONS and build_dataset_v1; data/questions/dataset_v1.jsonl; notes/DATASET.md')

# 20: three pair panels.
s=slide('Pair rankings improve, but the strongest model varies', 'Per-view accuracy: each distinct pair contributes its two presentation orders.', 'RESULTS')
for x,key,title,n in [(.89,'smaller_area','Choose the smaller area',46),(4.81,'lower_power','Choose lower power',44),(8.74,'higher_reach','Choose higher frequency',12)]:
    panel(s,x,2.0,3.7,4.08,title,f'{n//2} distinct pairs • {n} views',metric(key),hits(key),n,(.5,'always A or B = 50%'))
takeaway(s,'9B leads area ranking (40 of 46 views); 27B leads power ranking (31 of 44). Frequency has only six pairs.')
notes(s,'The full pair accuracy values are taken from the official local harness scorer. Per-view hits: area 22,35,40,36 of 46; power 28,25,29,31 of 44; frequency 8,10,9,10 of 12. The dashed 50% reference is always choosing A or always choosing B because both orders are included. This is not independent random sampling: the two orders share designs, and different pairs can share candidates of a problem. No significance test is claimed, and different per-task maxima belong to different models. All percentage axes are 0–100%.')

# 21: both orders.
s=slide('Check the design ranking in both orders', 'A stronger test counts a pair correct only when both presentation orders are correct.', 'RESULTS')
table(s,['Model','Area: both correct\n23 pairs','Power: both correct\n22 pairs','Frequency: both correct\n6 pairs'],
      [[NAMES[i]]+[f"{D['models'][m]['pair_counts'][k]['both_correct']} of {D['models'][m]['pair_counts'][k]['pairs']}" for k in ['smaller_area','lower_power','higher_reach']] for i,m in enumerate(M)],
      .89,2.05,[1.5,3.35,3.35,3.35],2.45,size=21,head_size=16)
txt(s,'Same letter after swapping the power pair:',1.08,4.88,11.1,.42,size=22,bold=True)
txt(s,'0.8B: 16 of 22 pairs   •   4B: 19 of 22   •   9B: 11 of 22   •   27B: 3 of 22',1.08,5.48,11.1,.4,size=19)
takeaway(s,'Keeping the same letter after a swap guarantees at least one wrong view. Lower is better for that error count.')
notes(s,'Both-order correct counts are derived from matching AB/BA answers and requiring each to equal its label. Area:5,13,19,17 of23; power:6,3,9,14 of22; frequency:3,4,4,5 of6. Same-letter-both-orders is the saved scorer positional inconsistency measure; power counts are16,19,11,3 of22. A consistent switched-letter answer can still choose the wrong design in both orders, so consistency alone does not prove correctness. The both-orders criterion is therefore shown separately from the per-view accuracy. It is a stricter descriptive summary, not a new experiment.')

# 22: runtime.
s=slide('All four Condor evaluations completed', 'Observed run times include different hardware and setup costs; they are not a controlled speed ranking.', 'EXECUTION')
table(s,['Model','Job','GPU used','Job time\nincluding setup','Median model\ntime per request'],[
    [NAMES[i],D['models'][m]['manifest']['job_id'], 'RTX 3080 Ti' if i==0 else 'A100 80 GB',f"{D['models'][m]['manifest']['condor_execution_seconds']:.0f} s",f"{D['models'][m]['manifest']['median_model_latency_ms']:.0f} ms"] for i,m in enumerate(M)],
    .89,2.08,[1.5,.8,3.15,2.95,3.15],2.65,size=21,head_size=16)
txt(s,'Every model: 348 dataset requests answered, 5 smoke requests passed,\n725 test questions scored, zero HTTP errors and zero missing answers.',1.08,5.08,11.1,.9,size=23,bold=True)
takeaway(s,'All runs used the same dataset fingerprint and v1.0 checkpoints. Maximum reported input length: 7,706 tokens.')
notes(s,'Condor job IDs4.0,5.0,6.0,7.0 completed with return code0. Job execution seconds are125,187,204,400; this includes setup/smoke/inference but excludes idle queue wait. Median server-reported model times are109.05,201.35,243.7,668.3ms. These medians include all348 dataset requests, development plus test, but exclude the separate smoke file. GPU0.8B: euler RTX3080Ti; larger models: aristoteles A100-SXM4-80GB. Sum of client wall times over the348 requests:52.047,99.433,104.114,279.849s. Server model latency and client wall time are distinct measures. Peak sampled GPU memory is recorded for0.8B only (7036MiB); no comparable peak is available for the larger jobs. An earlier0.8B attempt3.0 failed before inference on AMD and was corrected by requiring CUDA. No dollar-cost estimate is measured.', 'results/kev/*/manifest.json; condor/logs/4.0.kev-0.8b.log; condor/logs/5.0.kev-4b.log; condor/logs/6.0.kev-9b.log; condor/logs/7.0.kev-27b.log')

# 23: supported uses.
s=slide('What the evidence supports—and what needs work', 'These are observations on dataset v1, not a guarantee for unseen design families.', 'INTERPRETATION')
card(s,.89,2.0,5.65,3.85,'Promising in this test','27B: source-compilation screening.\n27B: failure reasons when failure is known.\n27B: coarse area levels.\n9B: pairwise area ranking.\n27B: pairwise power ranking.',C['navy'],20)
card(s,6.79,2.0,5.65,3.85,'Next experiments','Choose binary thresholds on development data.\nExpand synthesis failures and rare classes.\nAdd designs that miss the timing target.\nCheck power and timing across design groups.\nCompare other models under the same protocol.',C['teal'],18)
takeaway(s,'Keep tool verification in the loop. Reliable synthesis-failure screening is not established by these scores.')
notes(s,'This slide distinguishes measured strengths from proposed follow-up work. The test-set probabilities suggest that the fixed0.5 threshold may hide some ranking signal, but any alternative threshold must be selected on development only, frozen and assessed on new held-out data rather than optimised on this test. Expanded failure cases should include compiler-accepted synthesis failures. The current test has no negative-slack cases, no level5 power examples, and no level5 reachable-frequency examples. No confidence interval or paired significance test is claimed here. Jev, base Qwen, commercial LLMs, renamed/comment controls, energy per operation and post-route or measured-silicon validation have no results in this completed evaluation.')

# 24: appendix divider.
s=slide('',section='APPENDIX',bg=C['navy'])
for sh in list(s.shapes): sh._element.getparent().remove(sh._element)
TEXT_BOXES[:]=[b for b in TEXT_BOXES if b['slide']!=len(prs.slides)]
txt(s,'THE DETAIL BEHIND THE NUMBERS',.89,2.0,11.2,.4,size=16,color=C['cyan'],bold=True)
txt(s,'Full results and\ncoverage limits',.89,2.7,10.9,1.65,size=43,color=C['white'],bold=True)
txt(s,'Exact counts, probability metrics, level errors and reproducible evidence',.89,5.05,11.2,.75,size=23,color='BFDDEC')
SLIDE_META[-1]['title']='Full results and coverage limits'
notes(s,'Appendix for review and reproduction. The following tables retain all primary question metrics, the class splits, reason coverage, probability metrics, expected-level error and pair-order checks. There is no single overall cross-task accuracy because tasks have different meanings, supports and class distributions.')

# 25: all primary metrics.
LABELS={'compiles':'Compiles — pooled','synthesizes':'Synthesizes — pooled','reason':'Failure reason — known failure','area':'Area — exact level','power':'Power — exact level','reach':'Clock frequency — exact level','slack':'Timing margin — exact level','smaller_area':'Pair: smaller area','lower_power':'Pair: lower power','higher_reach':'Pair: higher frequency'}
ORDER=list(LABELS)
s=slide('Every primary test score', 'Correct answers (%) • each question has equal weight within its row • higher is better.', 'APPENDIX')
table(s,['Task','Questions']+SHORT,[[LABELS[k],D['models'][M[0]]['table'][k]['n']]+[pct(v) for v in metric(k)] for k in ORDER],
      .89,2.04,[4.1,1.25,1.55,1.55,1.55,1.55],4.5,size=16,head_size=15)
txt(s,'Pooled yes/no scores conceal class imbalance. Use the successful/failing split on the next slide.',.95,6.67,11.4,.25,size=12.5,color=C['muted'])
notes(s,'This is the complete primary accuracy table from the saved scorer, reproduced from raw answers. The KEV model size is indicated by the four numeric columns; all are KEV v1.0. Yes/no scores use threshold0.5. Reason accuracy is conditional on the prompt supplying failure. Area/power/frequency/slack require exact ordinal level. Pair scores count both presentation orders as separate views; area46views, power44views, frequency12views. No row is a general combined model score. For probability, class and level-error metrics see later appendix slides.')

# 26: explicit splits all models.
s=slide('Yes/no checks: every class, count and weighting', '“Correct” below is measured within the actual tool-outcome group.', 'APPENDIX')
rows=[]
for key,label in [('compiles','Compilation'),('synthesizes','Synthesis')]:
    for outcome,title in [('pass','Tool succeeds'),('fail','Tool fails')]:
        n=D['models'][M[0]]['binary'][key]['n_'+outcome]
        rows.append([label+': '+('successes' if outcome=='pass' else 'failures'),str(n)]+[f"{D['models'][m]['binary'][key][outcome+'_correct']} ({pct(D['models'][m]['binary'][key][outcome+'_rate'])})" for m in M])
    rows.append([label+': balanced','—']+[pct(v) for v in metric(key,'balanced_accuracy')])
    rows.append([label+': “yes” answers','123']+[str(D['models'][m]['binary'][key]['yes_answers']) for m in M])
table(s,['Actual tool group','Designs']+SHORT,rows,.89,2.05,[3.2,1.1,1.8,1.8,1.8,1.8],4.0,size=17,head_size=15)
txt(s,'Group rows: number correct, then the percentage of that group. “Yes” rows count all 123 designs.\nBalanced accuracy averages success recognition and failure detection with 50% weight each.',1.05,6.17,11.1,.7,size=16)
notes(s,'The rows show all four binary confusion-matrix supports with correct counts. Wrong counts are n minus correct. Compilation supports114pass/9fail; synthesis99pass/24fail. Balanced accuracy is the arithmetic mean of the two within-group percentages, not weighted by123designs. KEV0.8B synthesis always predicts yes. KEV4B compilation always predicts yes. The source answers are all present and successfully scored; there is no missing-answer adjustment.'
      ' The “yes” rows count the answers with a yes probability of 0.5 or more among all 123 test designs: correct answers on successes plus wrong answers on failures.')

# Appendix: every yes probability behind the yes/no scores.
s = slide('Most failing designs get the same answer as successes', 'Each dot is one held-out design. A yes probability of 0.5 or more counts as “yes”.', 'APPENDIX')
strip_panel(s, .89, 1.99, 5.65, 4.08, 'Does it synthesize?', 'synthesizes', ['tool synthesizes', 'tool fails'])
strip_panel(s, 6.79, 1.99, 5.65, 4.08, 'Does it compile?', 'compiles', ['vlog accepts', 'vlog rejects'])
# The takeaway claims that no other answer set is right for most designs of both groups.
RIGHT_BOTH = [(m, key) for m in M for key in ['compiles', 'synthesizes']
              if D['models'][m]['binary'][key]['pass_rate'] > .5 and D['models'][m]['binary'][key]['fail_rate'] > .5]
assert RIGHT_BOTH == [('kev-27b', 'compiles')], RIGHT_BOTH
# The title claims that, in every other answer set, most failures get the answer that most successes get.
for m in M:
    for key in ['compiles', 'synthesizes']:
        b = D['models'][m]['binary'][key]
        same = b['fail_wrong'] if b['pass_rate'] > .5 else b['fail_correct']
        assert (same / b['n_fail'] > .5) == ((m, key) != ('kev-27b', 'compiles')), (m, key)
b27 = D['models']['kev-27b']['binary']['compiles']
takeaway(s, f'Only KEV-27B’s compile answers are right in both groups: {b27["pass_correct"]} of {b27["n_pass"]} and {b27["fail_correct"]} of {b27["n_fail"]}.')
COMP_27 = {g: [p for _, p in D['models']['kev-27b']['noul']['compiles'][g]] for g in ['pass', 'fail']}
SUB_27 = D['models']['kev-27b']['synthesis_failure_subsets']
notes(s, 'Each dot is the saved noul answer (the yes probability) of one held-out basic request. In each model row, the upper teal dots are the designs that the tool accepts and the lower orange dots are the designs that the tool rejects. '
    'The small vertical spread in each lane comes from a fixed hash of the question ID and only reduces overlap; it carries no data. A probability of 0.5 or more counts as “yes”, as in the saved scorer. '
    f'“Yes” answers of 123, synthesis: {", ".join(map(str, YES["synthesizes"]))}; compilation: {", ".join(map(str, YES["compiles"]))} (0.8B, 4B, 9B, 27B). '
    f'KEV-0.8B gives every design a compile probability between {min(COMP_08):.2f} and {max(COMP_08):.2f} and a synthesis probability between {min(SYN_P["kev-0.8b"] + SYN_F["kev-0.8b"]):.2f} and {max(SYN_P["kev-0.8b"] + SYN_F["kev-0.8b"]):.2f}, for successes and failures alike '
    f'(AUROC {D["models"]["kev-0.8b"]["table"]["compiles"]["auroc"]:.3f} and {D["models"]["kev-0.8b"]["table"]["synthesizes"]["auroc"]:.3f}; 0.5 is chance). '
    f'KEV-27B compile probabilities: successes {min(COMP_27["pass"]):.2f} to {max(COMP_27["pass"]):.2f}, failures {min(COMP_27["fail"]):.2f} to {max(COMP_27["fail"]):.2f} (AUROC {D["models"]["kev-27b"]["table"]["compiles"]["auroc"]:.3f}). '
    f'KEV-27B synthesis probabilities: the {SUB_27["compiler_rejects"]["n"]} failures that also fail vlog rank below the successes (AUROC {SUB_27["compiler_rejects"]["auroc_vs_successes"]:.3f}), but only {SUB_27["compiler_rejects"]["correct"]} fall below 0.5; '
    f'the {SUB_27["compiler_accepts"]["n"]} failures that pass vlog get {SUB_27["compiler_accepts"]["min_p"]:.2f} to {SUB_27["compiler_accepts"]["max_p"]:.2f} (AUROC {SUB_27["compiler_accepts"]["auroc_vs_successes"]:.3f}). '
    'The Kev README names no RTL training data. A search of the Kev suites (evals/) for Verilog terms finds only the C++ logging macro VLOG and the word “vlog” in a text about videos. '
    'Kev-27B also trained on C++ code-review questions (devtools-v1) and starts from a post-trained Qwen model whose training data is not published.',
    'results/kev/*/dataset_v1__*.jsonl; presentations/src/analyse_kev.py; github.com/jaredpalmer/kev README and evals/ (commit 5e42a7a)')

# 27: reason classes.
REASON_NAMES={'multiple_drivers':'More than one driver','illegal_code':'Illegal code','unsupported_construct':'Unsupported construct','timing_control':'Unsupported timing control','simulation_statement':'Simulation-only statement','async_reset':'Asynchronous reset','mixed_assignments':'Mixed assignment styles','instance_mismatch':'Module-instance mismatch'}
s=slide('Failure-reason coverage is small and uneven', 'Correct reason counts by actual class • the prompt supplies that failure occurred.', 'APPENDIX')
rows=[[REASON_NAMES[k],n]+[D['models'][m]['reason_by_class'][k]['correct'] for m in M] for k,n in D['test_reason_counts'].items()]
rows.append(['Equal weight per observed class','Accuracy']+[pct(v) for v in metric('reason','balanced_accuracy')])
table(s,['Actual failure class','Cases']+SHORT,rows,.89,2.04,[4.1,1.25,1.55,1.55,1.55,1.55],4.13,size=16,head_size=15)
txt(s,'No true event-list or unbounded-loop cases occur in this held-out subset.\nThese two offered classes therefore have no measured class recall.',1.05,6.35,11.1,.55,size=15.5)
notes(s,'The eight observed reason classes and supports are drawn from the test input. They total16. The offered option list has10classes, so event_list and unbounded_loop have zero true cases and are excluded from the saved scorer balanced accuracy. Correct class counts for every model are recomputed. KEV27B sole mistake is multiple_drivers→async_reset. Five simulation-only examples and three timing-control examples give those classes more influence in pooled accuracy than one-case classes. Balanced accuracy instead weights the eight observed class recalls equally. Do not interpret this as broad10-class validation.', 'notes/synth_reason.md; data/questions/dataset_v1.jsonl; presentations/data/verified_metrics.json')

# 28: probability metrics.
s=slide('Probability ranking and binary decisions are different', 'A model can rank cases reasonably while making poor decisions at the fixed 0.5 threshold.', 'APPENDIX')
rows=[]
for key,label in [('compiles','Compilation'),('synthesizes','Synthesis')]:
    rows.append([label+': ranking score (AUROC)']+[f'{v:.3f}' for v in metric(key,'auroc')])
    rows.append([label+': probability error (Brier)']+[f'{v:.3f}' for v in metric(key,'brier')])
table(s,['Probability metric']+SHORT,rows,.89,2.05,[4.55,1.75,1.75,1.75,1.75],2.35,size=17,head_size=16)
card(s,.89,4.75,5.65,1.65,'Ranking score: higher is better','Chance = 0.5; perfect ordering = 1.\nCompares yes probabilities for successes and failures.',body_size=16)
card(s,6.79,4.75,5.65,1.65,'Probability error: lower is better','Mean squared error of the yes probability.\nZero is perfect; confidently wrong answers cost more.',body_size=16)
notes(s,'AUROC is the probability that a randomly chosen actual success has a higher predicted yes probability than a randomly chosen actual failure, counting ties as half. It ranges0–1;0.5 corresponds to chance ordering. It is independent of the chosen decision threshold and is not failure recall. Brier score is mean((p_yes-y)^2),0–1with lower better. It measures probability error, not calibration alone. These metrics use all123binary test designs. Synthesis AUROC improves to0.814for27B although its thresholded failure recall is16.7%. Any threshold tuning must use the development split, not these test labels. No recalibrated or alternative-threshold result has been run.')

# 29: performance secondary metrics.
s=slide('Performance levels: tolerance and expected-level error', 'Within one level: higher is better. Mean level error: lower is better.', 'APPENDIX')
rows=[]
for key,label in [('area','Area'),('power','Power'),('reach','Clock frequency'),('slack','Timing margin')]:
    rows.append([label,'Within one level']+[pct(v) for v in metric(key,'within_one')])
    rows.append([label,'Mean level error']+[f'{v:.3f}' for v in metric(key,'mae_levels')])
table(s,['Metric','Measure']+SHORT,rows,.89,2.05,[2.15,2.45,1.7375,1.7375,1.7375,1.7375],3.95,size=16,head_size=15)
txt(s,'Mean level error uses the probability-weighted expected level, not just the most likely level.\nUnequal physical bin widths mean that one level is not a fixed physical error.',1.05,6.24,11.1,.6,size=16)
notes(s,'The saved scorer computes within_one from the argmax chosen level versus the true integer level. Mean absolute level error (mae_levels) uses the probability-weighted expected level sum(k*p_k)/sum(p_k), then averages abs(expected-true). Therefore it is not the mean absolute error of the argmax level and need not track exact accuracy monotonically. Supports:92area,85power,92reachablefrequency,92slack. Since area and power bins cover decades, ordinal error cannot be read as fixed relative or absolute physical error. All values are retained from the scorer, verified against raw responses with1e-12floating-point tolerance.')

# 30: complete pair order data.
s=slide('Pair order: all correctness and letter-bias counts', 'Counts are per distinct pair. Each pair has exactly two model answers.', 'APPENDIX')
rows=[]
for key,label in [('smaller_area','Area'),('lower_power','Power'),('higher_reach','Frequency')]:
    n=D['models'][M[0]]['pair_counts'][key]['pairs']
    for stat,label2 in [('both_correct','Both orders correct'),('one_correct','One order correct'),('neither_correct','Neither order correct')]:
        rows.append([label+' — '+label2,str(n)]+[D['models'][m]['pair_counts'][key][stat] for m in M])
table(s,['Pair outcome','Pairs']+SHORT,rows,.89,2.05,[4.1,1.25,1.55,1.55,1.55,1.55],4.25,size=16,head_size=15)
txt(s,'With two opposite correct letters, “one order correct” is the same as keeping the same letter.\nA consistent letter swap can still be wrong in both orders.',1.05,6.44,11.1,.55,size=15)
notes(s,'For an A/B pair with opposite correct labels in AB and BA, each model response falls into0,1or2correct views. Both-order counts are read from paired saved answers, not estimated from percentages. The saved same_letter_both_orders rate equals the one-correct count divided by the number of distinct pairs for this balanced binary protocol. It is not itself an accuracy measure. Questions share design candidates and problem families, so naive independence across views is inappropriate. All row triplets sum to the labelled number of pairs.')

# 31: corpus composition and omitted work.
GROUP_NAMES={'metrex':'MetRex: varied sizes','mutant':'Injected synthesis errors','natural':'Natural language-model candidates','noncompile':'Injected compilation errors','reference':'Human reference designs','revold':'Older REvolution candidates','revolution':'REvolution v0 designs','trap':'Accepted edits and rejected twins'}
s=slide('Dataset composition and missing coverage', 'The retained designs cover several sources; this is not a balanced population sample.', 'APPENDIX')
table(s,['Design group','All retained','Held-out designs','Held-out synthesis failures'],
      [[GROUP_NAMES[g],v['all_designs'],v['test_designs'],v['test_synthesis_failures']] for g,v in D['groups'].items()],
      .89,2.05,[5.55,1.7,1.9,2.4],3.83,size=16,head_size=15)
txt(s,'Not measured here: energy per operation, functional equivalence, post-route timing,\nmeasured chips, other model families, or the older plan’s rename/comment controls.',1.05,6.14,11.1,.72,size=17)
notes(s,'The group counts are based on retained basic design records only, not requests or pair views. All retained counts total151 and held-out counts123. The original plan159had8excludeddesigns:4mutantswithanunintendedclass,2traptwins that did not synthesize as required, and2noncompile designs whose DC class was not intended. Twelve retained performance cases across both splits have no usable power value, leaving98powerquestions overall and85test. The corpus also skips comparisons with incompatible testbenches or less than10%metric differences. Some source groups contain near variants; base_design_id grouping prevents cross-split leakage as checked in analyse_kev.py. Dataset v1 supersedes older planning sizes such as500perquestiontype. No planned result is presented as completed.', 'data/questions/dataset_v1.card.md; notes/DATASET.md; harness/build.py; presentations/data/verified_metrics.json')

# 32: provenance.
s=slide('Evidence needed to reproduce this deck', 'Every number is tied to the frozen dataset and saved responses; no model reruns were needed.', 'APPENDIX')
table(s,['Model','Condor job','Completed on 15 October','Checkpoint revision (prefix)'],
      [[NAMES[i],D['models'][m]['manifest']['job_id'],D['models'][m]['manifest']['completion_time'].split()[1],D['models'][m]['manifest']['model_revision'][:12]] for i,m in enumerate(M)],
      .89,2.05,[1.65,1.4,4.0,4.5],2.5,size=19,head_size=16)
txt(s,'Inputs and evidence',1.05,4.86,11.1,.4,size=22,bold=True)
txt(s,'Questions: data/questions/dataset_v1.jsonl\nAnswers, scores and manifests: results/kev/\nVerified tables and rebuild scripts: presentations/data/ and presentations/src/',
    1.05,5.43,11.1,1.0,size=18)
notes(s,'Full provenance:\n'+json.dumps({m:{'job':D['models'][m]['manifest']['job_id'],'revision':D['models'][m]['manifest']['model_revision'],'answers_file':D['models'][m]['answers_file'],'answers_sha256':D['models'][m]['answers_sha256'],'dataset_sha256':D['dataset_sha256']} for m in M},indent=2)+'\nAll times are Europe/Paris (CEST) as recorded in the manifests. The build verifies unique request IDs, successful HTTP responses, matching question keys, no cross-split base IDs, input fingerprints, and saved metrics versus freshly recomputed metrics. Numerical floating-point differences up to1e-12are accepted. Theme and slide masters are inherited from the user-supplied reference. This deck is derived from the completed results only.')


def validate_and_save():
    issues=[]
    for i,s in enumerate(prs.slides,1):
        assert s.has_notes_slide and len(s.notes_slide.notes_text_frame.text)>200
        for sh in s.shapes:
            if sh.left < -100 or sh.top < -100 or sh.left+sh.width > prs.slide_width+100 or sh.top+sh.height > prs.slide_height+100:
                issues.append(f'Slide {i}: out-of-bounds shape {sh.name}')
    assert not issues, issues
    DEST.mkdir(exist_ok=True)
    if OUTPUT.exists():
        bak=DEST/'backups';bak.mkdir(exist_ok=True)
        shutil.copy2(OUTPUT,bak/(OUTPUT.stem+'_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.pptx'))
    prs.save(OUTPUT)
    # Load the saved package, ensuring it can be read without PowerPoint repair.
    check=Presentation(OUTPUT)
    assert len(check.slides)==33
    assert len(check.slides)==len(SLIDE_META)
    assert all(s.has_notes_slide for s in check.slides)
    (DEST/'qa/slide_manifest.json').write_text(json.dumps(SLIDE_META,indent=2)+'\n')
    (DEST/'qa/text_boxes.json').write_text(json.dumps(TEXT_BOXES,indent=2)+'\n')
    print(f'Saved {OUTPUT.relative_to(ROOT)}: {len(check.slides)} slides; all shapes in bounds; notes on every slide.')


if __name__=='__main__':
    validate_and_save()
