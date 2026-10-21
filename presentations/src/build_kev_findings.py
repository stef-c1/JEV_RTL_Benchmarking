"""Build the KEV findings deck: an editable PowerPoint in the TinyTapeOut theme.

Run from the project root:
  /users/micas/scuycken/no_backup/miniconda3/envs/mx_env/bin/python presentations/src/build_kev_findings.py
  Add --cached to reuse presentations/data/verified_metrics*.json instead of recomputing them.

Structure (MV, 2026-10-21: the deck of 19 October was chaotic): one slide per question, and on each
result slide the two experiments side by side: no examples (15 October) and three solved examples
(16 October). analyse_kev.py recomputes and verifies every number from the saved answers. The
claims in the titles, the tables, and the takeaways are checked with assertions.
"""
from __future__ import annotations

import hashlib
import json
import math
import re
import shutil
import sys
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
from analyse_kev import ROOT, MODELS  # noqa: E402

DEST = ROOT / 'presentations'
OUTPUT = DEST / 'KEV_RTL_findings_20261015.pptx'
TEMPLATE = ROOT / 'pwp_example/TinyTapeOut 2026 jullie feedback.pptx'
PYTHON = '/users/micas/scuycken/no_backup/miniconda3/envs/mx_env/bin/python'
if '--cached' in sys.argv:
    D = json.loads((DEST / 'data/verified_metrics.json').read_text())
    D2 = json.loads((DEST / 'data/verified_metrics_fewshot3.json').read_text())
else:
    from analyse_kev import analyse, save, analyse_fewshot, save_fewshot
    D = analyse()
    save(D)
    D2 = analyse_fewshot(D)
    save_fewshot(D2)

M = [x[0] for x in MODELS]
SHORT = ['0.8B', '4B', '9B', '27B']
NAMES = ['KEV-' + s for s in SHORT]
DM, D2M, CMP = D['models'], D2['models'], D2['compare']
C = dict(navy='00407A', bg='F4F8FA', pale='E3EFF5', cyan='52BDEC', teal='1D8DB0', blue='0B5A93',
         muted='4A5A66', grid='C8D3DA', gray='B5BFC3', white='FFFFFF', orange='E97132')
E1, E2 = C['cyan'], C['navy']          # bar colors: no examples, three examples
LEG1, LEG2 = 'no examples', 'three examples'
FOOTER = 'KEV v1.0  •  Dataset v1  •  runs of 15 and 16 October 2026'
FONT = {False: '/usr/share/fonts/liberation-sans/LiberationSans-Regular.ttf',
        True: '/usr/share/fonts/liberation-sans/LiberationSans-Bold.ttf'}

prs = Presentation(TEMPLATE)
for sid in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sid.rId)
    prs.slides._sldIdLst.remove(sid)
prs.core_properties.title = 'KEV on RTL: what the four KEV sizes can tell about an RTL design'
prs.core_properties.subject = 'Dataset v1; KEV v1.0 0.8B, 4B, 9B, 27B; 725 test questions; without and with three solved examples'
prs.core_properties.author = 'JEV_RTL study'
prs.core_properties.keywords = 'RTL, KEV, synthesis, compilation, area, power, timing'
TEXT_BOXES = []
SLIDE_META = []
WARN = []


# ----------------------------------------------------------------------------------------------
# Small helpers
# ----------------------------------------------------------------------------------------------
def rgb(c):
    return RGBColor.from_string(c)


def pct(value, digits=1):
    return str(Decimal(str(value * 100)).quantize(Decimal('1.' + '0' * digits), rounding=ROUND_HALF_UP)) + '%'


def width_in(text, size, bold=False):
    """The width of the widest line of a text in inches (Liberation Sans has the Arial metrics)."""
    font = ImageFont.truetype(FONT[bold], int(size * 10))
    return max(font.getlength(line) for line in str(text).split('\n')) / 10 / 72


def metric(key, field='accuracy', d=None):
    return [(d or D)['models'][m]['table'][key][field] for m in M]


def hits(key, d=None):
    return [round(t * (d or D)['models'][m]['table'][key]['n']) for t, m in zip(metric(key, d=d), M)]


def n_of(key):
    n = DM[M[0]]['table'][key]['n']
    assert all(d['models'][m]['table'][key]['n'] == n for d in (D, D2) for m in M)
    return n


def ref_level(key):
    """The constant level that is right most often in the test labels: (level, count, n)."""
    t = DM[M[0]]['table'][key]
    assert all(d['models'][m]['table'][key]['truth_counts'] == t['truth_counts'] for d in (D, D2) for m in M)
    counts = {int(k): v for k, v in t['truth_counts'].items()}
    level = max(counts, key=counts.get)
    return level, counts[level], t['n']


def one_answer(m, key, d):
    """The answer of a model if it gives the same answer to every question of the key, else None."""
    pred = d['models'][m]['table'][key]['pred_counts']
    return next(iter(pred)) if len(pred) == 1 else None


def level_counts(m, key, d, field='pred_counts'):
    t = d['models'][m]['table'][key]
    return [t[field].get(str(i), 0) for i in range(1, 6)]


def pval(p):
    return 'p < 0.001' if p < .001 else f'p = {p:.2g}'


def change(a, b):
    return f'{a} → {b}'


# ----------------------------------------------------------------------------------------------
# Drawing helpers
# ----------------------------------------------------------------------------------------------
def box(s, x, y, w, h, fill=None, line=None, radius=False):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    if fill:
        shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = rgb(line); shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    if radius:
        try:
            shape.adjustments[0] = 0.06
        except (IndexError, TypeError):
            pass
    return shape


def txt(s, text, x, y, w, h, size=18, color=None, bold=False, align=PP_ALIGN.LEFT,
        valign=MSO_ANCHOR.TOP, margin=0, record=True):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, line_ in enumerate(str(text).split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.09
        r = p.add_run(); r.text = line_
        r.font.name = 'Arial'; r.font.size = Pt(size); r.font.bold = bold
        r.font.color.rgb = rgb(color or C['navy'])
    if record:
        TEXT_BOXES.append(dict(slide=len(prs.slides), text=str(text), x=x, y=y, w=w, h=h, size=size))
    return shape


def line(s, x1, y1, x2, y2, color=None, width=1):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = rgb(color or C['grid']); sh.line.width = Pt(width)
    return sh


def arrow_line(s, x1, y1, x2, y2, color=None, width=2):
    """A straight connector with an arrowhead at its end."""
    sh = line(s, x1, y1, x2, y2, color or C['teal'], width)
    ln = sh.line._get_or_add_ln()
    tail = OxmlElement('a:tailEnd'); tail.set('type', 'triangle'); tail.set('w', 'med'); tail.set('len', 'med')
    ln.append(tail)
    return sh


def notes(s, content, source=None):
    text = (content.strip() + '\n\nEvidence and reproducibility\n' +
            (source or 'data/questions/dataset_v1.jsonl; data/questions/dataset_v1_fewshot3.jsonl; results/kev/*/; harness/score.py; harness/compare.py') +
            '\nAll results use the test split: 123 designs, 725 questions. No examples = experiment 1 (Condor jobs 4.0 to 7.0, 15 October 2026); '
            'three examples = experiment 2 (Condor jobs 8.0 to 11.0, 16 October 2026).' +
            '\nDataset SHA-256: ' + D['dataset_sha256'] + ' (dataset_v1.jsonl); ' + D2['dataset_sha256'] + ' (dataset_v1_fewshot3.jsonl).' +
            '\nVerified calculations: presentations/data/verified_metrics.json, verified_metrics_fewshot3.json, compare_fewshot3.csv (presentations/src/analyse_kev.py).' +
            '\nRebuild: ' + PYTHON + ' presentations/src/build_kev_findings.py' +
            '\nVisual style: pwp_example/TinyTapeOut 2026 jullie feedback.pptx; theme and slide masters retained. All charts and tables are editable PowerPoint objects.')
    ns = s.notes_slide
    if ns.notes_text_frame is None:
        # The supplied reference has no body placeholder in its notes master.
        sp = OxmlElement('p:sp')
        nv = OxmlElement('p:nvSpPr')
        prop = OxmlElement('p:cNvPr'); prop.set('id', '100'); prop.set('name', 'Speaker notes')
        nv.append(prop); nv.append(OxmlElement('p:cNvSpPr'))
        nvp = OxmlElement('p:nvPr'); ph = OxmlElement('p:ph'); ph.set('type', 'body'); ph.set('idx', '13')
        nvp.append(ph); nv.append(nvp); sp.append(nv); sp.append(OxmlElement('p:spPr'))
        body = OxmlElement('p:txBody')
        for name in ['a:bodyPr', 'a:lstStyle', 'a:p']:
            body.append(OxmlElement(name))
        sp.append(body); ns._element.cSld.spTree.append(sp)
    ns.notes_text_frame.text = text


def bar_legend(s, x, y, size=12):
    """The legend of the bar colors: light = no examples, dark = three examples."""
    box(s, x, y + .06, .2, .2, fill=E1)
    txt(s, LEG1, x + .28, y, 1.2, .3, size=size, color=C['muted'])
    box(s, x + 1.5, y + .06, .2, .2, fill=E2)
    txt(s, LEG2, x + 1.78, y, 1.4, .3, size=size, color=C['muted'])


def slide(title, subtitle='', section='RESULTS', bg=None, legend=False):
    s = prs.slides.add_slide(prs.slide_layouts[-1])
    for sh in list(s.shapes):
        sh._element.getparent().remove(sh._element)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg or C['bg'])
    if width_in(title, 31, True) > 11.5:
        WARN.append(f'title too wide ({width_in(title, 31, True):.2f} in): {title}')
    txt(s, title, .89, .65, 11.55, .68, size=31, bold=True)
    if subtitle:
        wsub = 8.35 if legend else 11.55
        if width_in(subtitle, 16) > wsub:
            WARN.append(f'subtitle too wide ({width_in(subtitle, 16):.2f} in, limit {wsub}): {subtitle}')
        txt(s, subtitle, .89, 1.38, wsub, .48, size=16, color=C['muted'])
    if legend:
        bar_legend(s, 9.35, 1.42)
    txt(s, FOOTER + '  |  ' + section, .89, 7.04, 10.65, .22, size=10.5, color=C['muted'])
    txt(s, str(len(prs.slides)), 11.85, 7.02, .58, .25, size=11, color=C['muted'], align=PP_ALIGN.RIGHT)
    SLIDE_META.append({'number': len(prs.slides), 'title': title, 'section': section})
    return s


def card(s, x, y, w, h, title, body, color=None, body_size=18, title_size=None):
    box(s, x, y, w, h, fill=C['pale'], radius=True)
    box(s, x, y, w, .045, fill=color or C['teal'])
    compact = h < 2.0
    txt(s, title, x + .25, y + (.14 if compact else .25), w - .5, .4 if compact else .72,
        size=title_size or (17 if compact else 21), bold=True)
    txt(s, body, x + .25, y + (.55 if compact else 1.05), w - .5, h - (.62 if compact else 1.2), size=body_size)


def takeaway(s, text, y=6.25, size=15):
    box(s, .89, y, 11.55, .66, fill=C['pale'], radius=True)
    txt(s, text, 1.08, y + .06, 11.17, .54, size=size, bold=True, valign=MSO_ANCHOR.MIDDLE)


def table(s, headers, rows, x, y, widths, h, size=16, head_size=14, row_header=False, head_h=.62, left_cols=(0,), nowrap=False):
    t = s.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y),
                           Inches(sum(widths)), Inches(h)).table
    for j, width in enumerate(widths):
        t.columns[j].width = Inches(width)
    t.rows[0].height = Inches(head_h)
    for i in range(1, len(rows) + 1):
        t.rows[i].height = Inches((h - head_h) / len(rows))
    for i, row in enumerate([headers] + rows):
        for j, value in enumerate(row):
            if nowrap and i > 0 and width_in(value, size, row_header and j == 0) > widths[j] - .22:
                WARN.append(f'table cell too wide ({width_in(value, size, row_header and j == 0):.2f} in, column {widths[j]}): {value}')
            cell = t.cell(i, j)
            cell.margin_left = Inches(.12); cell.margin_right = Inches(.10)
            cell.margin_top = Inches(.05); cell.margin_bottom = Inches(.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb = rgb(C['navy'] if i == 0 else (C['pale'] if i % 2 else C['white']))
            cell.text = str(value)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT if j in left_cols else PP_ALIGN.CENTER
                p.space_before = Pt(0); p.space_after = Pt(0); p.line_spacing = 1.03
                for r in p.runs:
                    r.font.name = 'Arial'; r.font.size = Pt(head_size if i == 0 else size)
                    r.font.bold = i == 0 or (row_header and j == 0)
                    r.font.color.rgb = rgb(C['white'] if i == 0 else C['navy'])
            for edge in ['lnL', 'lnR', 'lnT', 'lnB']:
                border = OxmlElement('a:' + edge)
                border.append(OxmlElement('a:noFill'))
                cell._tc.get_or_add_tcPr().append(border)
    return t


def gpanel(s, x, y, w, h, title, desc, vals, labels, reference=None, label_w=1.35, label_size=11):
    """A white panel with two bars per model: no examples (light) and three examples (dark).
    vals = (values without examples, values with examples), fractions of 1; labels = the same two
    lists with the one-line label of each bar; reference = (fraction or list of fractions, legend)."""
    box(s, x, y, w, h, fill=C['white'], radius=True)
    txt(s, title, x + .22, y + .16, w - .44, .4, size=18, bold=True)
    txt(s, desc, x + .22, y + .58, w - .44, .3, size=12.5, color=C['muted'])
    px, pw = x + .93, w - .93 - .3 - label_w
    ry, step, bh, off = y + 1.0, .68, .22, .27
    top, bottom = ry - .06, ry + 3 * step + off + bh + .05
    for tick in [0, 50, 100]:
        xx = px + pw * tick / 100
        line(s, xx, top, xx, bottom, color=C['grid'], width=.65)
        txt(s, str(tick), xx - .22, bottom + .02, .44, .22, size=11, color=C['muted'], align=PP_ALIGN.CENTER)
    if reference:
        refs = reference[0] if isinstance(reference[0], (list, tuple)) else [reference[0]]
        for ref in refs:
            ln = line(s, px + pw * ref, top, px + pw * ref, bottom, C['muted'], 1.2)
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    for i in range(4):
        gy = ry + i * step
        txt(s, SHORT[i], x + .2, gy + .1, .66, .3, size=14)
        for j, col in enumerate((E1, E2)):
            v, lab = vals[j][i], labels[j][i]
            by = gy + j * off
            if v > 0:
                box(s, px, by, pw * v, bh, fill=col)
            if width_in(lab, label_size) > label_w:
                WARN.append(f'bar label too wide ({width_in(lab, label_size):.2f} in, limit {label_w}): {lab}')
            txt(s, lab, px + pw + .1, by - .03, label_w + .15, .3, size=label_size)
    if reference:
        several = isinstance(reference[0], (list, tuple))
        txt(s, ('Dashed lines: ' if several else 'Dashed line: ') + reference[1], x + .2, y + h - .27,
            w - .4, .23, size=10.5, color=C['muted'])


def count_labels(key, d, n=None):
    n = n or n_of(key)
    return [f'{pct(v)} · {h} of {n}' for v, h in zip(metric(key, d=d), hits(key, d))]


def level_panel(s, x, y, w, h, key, title, desc):
    level, count, n = ref_level(key)
    gpanel(s, x, y, w, h, title, desc, (metric(key), metric(key, d=D2)), (count_labels(key, D), count_labels(key, D2)),
           reference=(count / n, f'always level {level} = {pct(count / n)} ({count} of {n})'))


def yes_no_labels(key, d):
    bs = [d['models'][m]['binary'][key] for m in M]
    vals = metric(key, d=d)
    n = bs[0]['n_pass'] + bs[0]['n_fail']
    assert all(math.isclose(v, (b['pass_correct'] + b['fail_correct']) / n) for v, b in zip(vals, bs))
    return [f'{pct(v)} · {b["pass_correct"]} of {b["n_pass"]} yes · {b["fail_correct"]} of {b["n_fail"]} no' for v, b in zip(vals, bs)]


def yes_no_panel(s, x, y, w, h, key, title):
    b = DM[M[0]]['binary'][key]
    n = b['n_pass'] + b['n_fail']
    assert all((d['models'][m]['binary'][key]['n_pass'], d['models'][m]['binary'][key]['n_fail']) == (b['n_pass'], b['n_fail']) for d in (D, D2) for m in M)
    lo, hi = b['n_fail'] / n, b['n_pass'] / n
    gpanel(s, x, y, w, h, title, f'The tool says “yes” for {b["n_pass"]} of the {n} designs and “no” for {b["n_fail"]}',
           (metric(key), metric(key, d=D2)), (yes_no_labels(key, D), yes_no_labels(key, D2)),
           reference=([lo, hi], f'always “no” = {pct(lo)} · always “yes” = {pct(hi)}'), label_w=2.35)


def dist3(s, x, y, w, h, title, desc, truth, pred1, pred2):
    """Design counts at each of the five levels: the tool, the model without examples, the model with examples."""
    box(s, x, y, w, h, C['white'], radius=True)
    txt(s, title, x + .22, y + .16, w - .44, .4, size=18, bold=True)
    txt(s, desc, x + .22, y + .58, w - .44, .3, size=12.5, color=C['muted'])
    px, py, pw, ph = x + .63, y + 1.2, w - .98, 1.9
    for tick in [0, 25, 50, 75, 100]:
        yy = py + ph * (1 - tick / 100)
        line(s, px, yy, px + pw, yy, C['grid'], .65)
        txt(s, str(tick), x + .07, yy - .11, .45, .23, size=11, color=C['muted'], align=PP_ALIGN.RIGHT)
    slot = pw / 5
    for i, counts in enumerate(zip(truth, pred1, pred2)):
        xx = px + i * slot + .1
        for j, (v, col) in enumerate(zip(counts, (C['gray'], E1, E2))):
            bx = xx + j * .26
            if v:
                box(s, bx, py + ph * (1 - v / 100), .21, ph * v / 100, col)
            txt(s, str(v), bx - .07, py + ph * (1 - v / 100) - .24, .35, .22, size=10.5, align=PP_ALIGN.CENTER)
        txt(s, str(i + 1), xx + .18, py + ph + .08, .4, .25, size=13, align=PP_ALIGN.CENTER)
    txt(s, 'Level', x + .22, y + h - .5, .6, .28, size=12, color=C['muted'])
    lx = x + 1.0
    for col, label, wl in [(C['gray'], 'tool', .55), (E1, LEG1, 1.1), (E2, LEG2, 1.3)]:
        box(s, lx, y + h - .45, .18, .18, col)
        txt(s, label, lx + .25, y + h - .5, wl, .28, size=12)
        lx += .35 + wl


# The yes probabilities of every design (appendix).
OK_COLOR, FAIL_COLOR = C['teal'], C['orange']  # validated pair on white: CVD dE 17.3, contrast 3.83 and 3.06


def dot(s, cx, cy, d, color):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = rgb(color)
    sh.line.color.rgb = rgb(C['white']); sh.line.width = Pt(.6)
    return sh


def legend(s, x, y, w, items, size=12.5):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(.3))
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    for i, (color, label) in enumerate(items):
        for text, col in [(('      ' if i else '') + '●', color), (' ' + label, C['muted'])]:
            r = p.add_run(); r.text = text
            r.font.name = 'Arial'; r.font.size = Pt(size); r.font.color.rgb = rgb(col)
    TEXT_BOXES.append(dict(slide=len(prs.slides), text='   '.join(l for _, l in items), x=x, y=y, w=w, h=.3, size=size))


def jitter(qid):
    """Fixed vertical offset in [-1, 1] per question ID, so every build draws the same dots."""
    return int.from_bytes(hashlib.sha256(qid.encode()).digest()[:4], 'big') / 2**31 - 1


def strip_panel(s, x, y, w, h, title, key, labels, d):
    box(s, x, y, w, h, fill=C['white'], radius=True)
    txt(s, title, x + .22, y + .18, w - .44, .55, size=18, bold=True)
    n = d['models'][M[0]]['binary'][key]
    legend(s, x + .22, y + .67, w - .44, [(OK_COLOR, f'{labels[0]}: {n["n_pass"]}'), (FAIL_COLOR, f'{labels[1]}: {n["n_fail"]}')])
    px, pw = x + .93, w - 2.35
    ry, step = y + 1.18, .55
    for tick in [0, .5, 1]:
        xx = px + pw * tick
        ln = line(s, xx, ry - .08, xx, ry + 3 * step + .4, color=C['muted'] if tick == .5 else C['grid'], width=1.2 if tick == .5 else .65)
        if tick == .5:
            ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        txt(s, f'{tick:g}', xx - .22, y + 3.22, .44, .22, size=11, color=C['muted'], align=PP_ALIGN.CENTER)
    for i, m in enumerate(M):
        yy = ry + i * step
        txt(s, SHORT[i], x + .2, yy + .025, .66, .3, size=14)
        for lane, group, color in [(yy + .08, 'pass', OK_COLOR), (yy + .30, 'fail', FAIL_COLOR)]:
            for qid, p in d['models'][m]['noul'][key][group]:
                dot(s, px + pw * p, lane + .055 * jitter(qid), .08, color)
        txt(s, f'{d["models"][m]["binary"][key]["yes_answers"]} of 123\nare “yes”', px + pw + .13, yy - .035, 1.1, .47, size=12.5)
    txt(s, 'Yes probability of each design • dashed line: the 0.5 threshold', x + .2, y + h - .57, w - .4, .24, size=11, color=C['muted'])


# ----------------------------------------------------------------------------------------------
# Shared numbers, checked against the data
# ----------------------------------------------------------------------------------------------
BIN = {(key, e): [d['models'][m]['binary'][key] for m in M] for key in ['compiles', 'synthesizes'] for e, d in [(1, D), (2, D2)]}
YES1 = {k: [b['yes_answers'] for b in BIN[(k, 1)]] for k in ['compiles', 'synthesizes']}
YES2 = {k: [b['yes_answers'] for b in BIN[(k, 2)]] for k in ['compiles', 'synthesizes']}
N_PASS = {k: BIN[(k, 1)][0]['n_pass'] for k in YES1}
N_FAIL = {k: BIN[(k, 1)][0]['n_fail'] for k in YES1}
assert (N_PASS['compiles'], N_FAIL['compiles'], N_PASS['synthesizes'], N_FAIL['synthesizes']) == (114, 9, 99, 24)
S27_1, S27_2 = DM['kev-27b']['binary']['synthesizes'], D2M['kev-27b']['binary']['synthesizes']
C27_1, C27_2 = DM['kev-27b']['binary']['compiles'], D2M['kev-27b']['binary']['compiles']
H = {(k, 1): hits(k) for k in ['compiles', 'synthesizes', 'reason', 'area', 'power', 'reach', 'slack', 'smaller_area', 'lower_power', 'higher_reach']}
H.update({(k, 2): hits(k, D2) for k, _ in H})
QKEYS = ['compiles', 'synthesizes', 'reason', 'area', 'power', 'reach', 'slack', 'smaller_area', 'lower_power', 'higher_reach']
CONST_08 = [k for k in QKEYS if one_answer('kev-0.8b', k, D2)]
assert CONST_08 == ['synthesizes', 'area', 'power', 'slack', 'smaller_area', 'higher_reach']
SUB = {(k, e): [d['models'][m]['synthesis_failure_subsets'][k] for m in M] for k in ['compiler_accepts', 'compiler_rejects'] for e, d in [(1, D), (2, D2)]}
assert all(x['n'] == 15 for x in SUB[('compiler_accepts', 1)] + SUB[('compiler_accepts', 2)])
assert all(x['n'] == 9 for x in SUB[('compiler_rejects', 1)] + SUB[('compiler_rejects', 2)])
RL = {e: [d['models'][m]['reason_linked'] for m in M] for e, d in [(1, D), (2, D2)]}
PC = {e: {m: d['models'][m]['pair_counts'] for m in M} for e, d in [(1, D), (2, D2)]}
REASON_MAX = max(D['test_reason_counts'].values())
REASON_MODE = [k for k, v in D['test_reason_counts'].items() if v == REASON_MAX]
assert REASON_MODE == ['simulation_statement'] and sum(D['test_reason_counts'].values()) == 16
EX = D2['examples']
REASON_EX = [e['answer'] for e in EX['reason']['examples']]
assert REASON_EX == ['multiple_drivers', 'illegal_code', 'event_list']
SIG = [(m, k) for m in M for k in QKEYS if CMP[m]['keys'][k]['mcnemar_p'] < .05]
# The constant answers of the level questions, and the best count of any model in either experiment.
LEVEL_REF = {k: ref_level(k) for k in ['area', 'power', 'reach', 'slack']}
BEST = {k: max(H[(k, 1)] + H[(k, 2)]) for k in QKEYS}
assert BEST['area'] > LEVEL_REF['area'][1]
# Power: only KEV-4B without examples (39 of 85) is above always level 2 (37 of 85). Frequency and slack: no score above.
assert BEST['power'] == H[('power', 1)][1] == LEVEL_REF['power'][1] + 2 and BEST['reach'] < LEVEL_REF['reach'][1] and BEST['slack'] == LEVEL_REF['slack'][1]
assert BEST['synthesizes'] == S27_2['pass_correct'] + S27_2['fail_correct'] == N_PASS['synthesizes'] + 9
assert H[('compiles', 1)][3] == 122 and H[('compiles', 2)][3] == 117


# ----------------------------------------------------------------------------------------------
# 1. Cover
# ----------------------------------------------------------------------------------------------
s = slide('', section='OVERVIEW', bg=C['navy'])
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)
TEXT_BOXES[:] = [b for b in TEXT_BOXES if b['slide'] != 1]
box(s, 9.35, 1.45, 2.0, 2.0, line=C['teal'], radius=True)
txt(s, 'JEV_RTL STUDY  •  KEV v1.0', .89, 1.95, 10, .35, size=16, color=C['cyan'], bold=True)
txt(s, 'What can KEV tell us\nabout an RTL design?', .89, 2.55, 9.1, 1.62, size=43, color=C['white'], bold=True)
txt(s, 'Eight questions, four model sizes, 123 test designs,\nwithout and with three solved examples', .89, 4.55, 9.9, 1.0, size=23, color='BFDDEC')
txt(s, 'Runs of 15 October 2026 (no examples) and 16 October 2026 (three examples)  •  Deck of 21 October 2026',
    .89, 6.3, 11.4, .4, size=16, color='BFDDEC')
SLIDE_META[0]['title'] = 'What can KEV tell us about an RTL design?'
notes(s, 'This deck reports the completed KEV v1.0 runs of the four sizes KEV-0.8B, KEV-4B, KEV-9B, and KEV-27B (B = about a billion parameters) on dataset v1 of the JEV_RTL study. '
      'RTL means register-transfer level: source code that describes digital hardware. KEV is the open decision model of the Jev family: it returns typed answers only (a yes probability, one choice, or one level), no text. '
      'Dataset v1 has 151 designs; the test split has 123 designs and 725 questions, the development split 28 designs and 150 questions. '
      'Experiment 1 (15 October 2026) sends the questions without examples. Experiment 2 (16 October 2026) sends the same 725 test questions, one per request, each with three solved development designs before the design. '
      'Each result slide shows both experiments side by side: the light bar is the score without examples, the dark bar the score with three examples. '
      'All findings are specific to this dataset and tool flow. No Jev, base Qwen, or commercial LLM result is included.')

# ----------------------------------------------------------------------------------------------
# 2. The test in one picture
# ----------------------------------------------------------------------------------------------
s = slide('The test in one picture', 'The tools and KEV get the same RTL design. The score counts the KEV answers that agree with the tools.', 'OVERVIEW')
card(s, .89, 2.5, 2.7, 1.55, 'RTL design', 'Verilog source, top module, flow description. Power questions also get the testbench.', C['navy'], 13)
card(s, 4.45, 1.98, 4.0, 1.3, 'Reference tools', 'QuestaSim vlog, Design Compiler, PrimeTime give the correct answer.', C['teal'], 13)
card(s, 4.45, 3.27, 4.0, 1.3, 'KEV 0.8B, 4B, 9B, 27B', 'One typed answer per question, no text: a yes probability, a choice, or a level.', C['blue'], 13)
card(s, 9.3, 2.5, 3.14, 1.55, 'Score', 'Correct answers divided by questions. Reference: the score of one fixed answer for every design.', C['navy'], 13)
arrow_line(s, 3.59, 3.0, 4.45, 2.63)
arrow_line(s, 3.59, 3.55, 4.45, 3.92)
arrow_line(s, 8.45, 2.63, 9.3, 3.0)
arrow_line(s, 8.45, 3.92, 9.3, 3.55)
card(s, .89, 4.85, 3.7, 1.25, 'Four KEV sizes', 'KEV v1.0 with 0.8B, 4B, 9B, and 27B parameters. Each size answers every question.', body_size=13)
card(s, 4.81, 4.85, 3.7, 1.25, '123 test designs', 'From 151 designs of five sources. The 28 development designs serve only as examples.', body_size=13)
card(s, 8.74, 4.85, 3.7, 1.25, 'Two experiments', 'No examples (15 October). Three solved development examples before each question (16 October).', body_size=13)
takeaway(s, 'KEV never sees the tool logs. A tool run of each design gives the correct answer of every question about that design.')
notes(s, 'Each design gets one QuestaSim vlog -sv run (compile check) and one run of MV’s flow at 500 MHz: Design Compiler U-2022.12-SP2 with compile_ultra (synthesis, cell area, worst slack) and PrimeTime (power with the power testbench), NanGate 15 nm, typical corner, 0.8 V, 25 C. '
      'These runs give the correct answers. The model receives the state fields rtl (all source files, top module last), top, flow (a fixed text with the tool, library, and constraints), and for the power questions tb_power. '
      'The correct_answer and evidence fields stay local for scoring; the Condor client sends only the state, the model name, and the questions to /v1/systemone. '
      'The questions of one design go in up to three requests (basic, reason, performance) so that no question shows the answer of another question; a pair request holds two designs. '
      'The score of a question type is the share of its test questions with a correct answer (plain accuracy; MV, 2026-10-19). The reference is the score of one fixed answer for every design, computed from the test labels.',
      'notes/README.md; notes/DATASET.md; harness/build.py; condor/kev_client.py')

# ----------------------------------------------------------------------------------------------
# 3. The designs
# ----------------------------------------------------------------------------------------------
s = slide('The designs: 151 designs from five sources', '875 questions on 151 designs, both splits. An edit (mutant, trap, noncompile) counts for the source of its base design.', 'OVERVIEW')
pic = s.shapes.add_picture(str(DEST / 'dataset_v1_sources.png'), Inches(.89), Inches(2.0), width=Inches(11.55))
assert pic.height / 914400 < 4.2
GROUPS = D['groups']
assert sum(g['all_designs'] for g in GROUPS.values()) == 151 and sum(g['test_designs'] for g in GROUPS.values()) == 123
GROUP_TEXT = ', '.join(f"{g} {v['test_designs']} (of {v['all_designs']})" for g, v in GROUPS.items())
takeaway(s, f'Test split: {sum(g["test_designs"] for g in GROUPS.values())} designs, 725 questions. Development split: 28 designs, 150 questions, used only for the examples of experiment 2.')
notes(s, 'The figure (presentations/src/plot_dataset_sources.py) counts the designs and the questions of dataset v1 per source dataset, both splits together. '
      'REvolution ASP-DAC 2026: LLM-written candidates of the REvolution loop for RTLLM and VerilogEval problems; these designs also form the pairs. MetRex: a public set of designs of many sizes. '
      'RTLLM v2.0: public, verified designs and ChatGPT-written files. ARES: own medium designs (v0). VerilogEval: public reference designs. '
      'A script edit of a base design (a mutant with one injected synthesis error, a trap pair of an accepted edit and a rejected twin, or a noncompile edit with one removed semicolon or one undeclared name) counts for the source of its base design. '
      'The plan ran 159 designs; 8 are excluded by the dataset rules (an unintended error class, or a twin that did not synthesize), so 151 designs remain: 123 in the test split and 28 in the development split. The split is per base design, so no edit of a test design is in the development split. '
      f'Test designs per group: {GROUP_TEXT}. The appendix lists the groups with their synthesis failures.',
      'presentations/src/plot_dataset_sources.py; presentations/data/dataset_v1_sources.csv; data/questions/dataset_v1.card.md; notes/DATASET.md')

# ----------------------------------------------------------------------------------------------
# 4. The eight questions
# ----------------------------------------------------------------------------------------------
TT = D['test_truth_counts']
PAIRS = {k: PC[1][M[0]][k]['pairs'] for k in ['smaller_area', 'lower_power', 'higher_reach']}
assert PAIRS == {'smaller_area': 23, 'lower_power': 22, 'higher_reach': 6}
assert D['test_request_counts'] == {'basic': 123, 'performance': 92, 'reason': 16, 'pair': 50}
s = slide('The eight questions', 'Each question has a typed answer. The correct answer of every question comes from a tool run of the same design.', 'OVERVIEW')
table(s, ['Question', 'KEV answers with', 'Correct answer from', 'Test questions'], [
    ['Does it compile?', 'a yes probability', 'QuestaSim vlog -sv', f'{n_of("compiles")} designs: {TT["compiles"]["True"]} yes, {TT["compiles"]["False"]} no'],
    ['Does it synthesize?', 'a yes probability', 'Design Compiler, compile_ultra', f'{n_of("synthesizes")} designs: {TT["synthesizes"]["True"]} yes, {TT["synthesizes"]["False"]} no'],
    ['Why does synthesis fail?', 'one of ten reasons', 'the Design Compiler message', f'{n_of("reason")} designs; the prompt says that it fails'],
    ['Area level', 'one of five levels', 'Design Compiler, cell area', f'{n_of("area")} designs'],
    ['Power level', 'one of five levels', 'PrimeTime, with the testbench', f'{n_of("power")} designs'],
    ['Reachable clock frequency level', 'one of five levels', 'Design Compiler, worst slack', f'{n_of("reach")} designs'],
    ['Timing margin (slack) level', 'one of five levels', 'Design Compiler, worst slack', f'{n_of("slack")} designs'],
    ['Which of two designs has the smaller area, the lower power, the higher frequency?', 'A or B', 'the same tool runs', f'{PAIRS["smaller_area"]}, {PAIRS["lower_power"]}, {PAIRS["higher_reach"]} pairs, each asked in both orders'],
], .89, 1.98, [4.25, 1.9, 2.55, 2.85], 4.1, size=14, head_size=14, head_h=.45, left_cols=(0, 1, 2, 3))
takeaway(s, 'Compile and synthesize are two questions in one request; the power question of a design is in the same request as its area and timing questions.')
notes(s, 'The eight question types of dataset v1 (notes/README.md). Does it compile: can QuestaSim vlog -sv accept the source? Does it synthesize: can Design Compiler build a complete gate-level netlist (no error, no unresolved module, no unmapped register)? '
      f'In the test split {TT["synthesizes"]["True"]} designs compile and synthesize, {TT["compiles"]["True"] - TT["synthesizes"]["True"]} compile but fail synthesis, and {TT["compiles"]["False"]} fail both. '
      'Why does synthesis fail: the prompt states that Design Compiler does not synthesize the module and offers ten reasons; only designs that fail synthesis with messages of one class get this question (16 test designs; 8 further test designs with an injected compile error get no reason question). '
      'Area, power, reachable frequency, and slack: the model chooses one of five levels (next slide); only designs that synthesize get these questions (92 test designs; 85 of them have a usable power value). '
      'Pairs: two REvolution candidates of one problem, with the same flow and the same power testbench; a pair counts for a metric only if the two values differ by at least 10%. Each pair is asked in both orders (A/B and B/A), so 23 area pairs give 46 questions, 22 power pairs 44, and 6 frequency pairs 12. '
      'The yes/no answer of KEV is a probability; 0.5 or more counts as yes, as in the scorer. Test questions: 123 + 123 + 16 + 92 + 85 + 92 + 92 + 46 + 44 + 12 = 725.',
      'notes/README.md; notes/DATASET.md; notes/synth_yesno.md; notes/synth_reason.md; harness/build.py; data/questions/dataset_v1.jsonl')

# ----------------------------------------------------------------------------------------------
# 5. The five levels
# ----------------------------------------------------------------------------------------------
s = slide('The five levels of area, power, and timing', 'The model chooses a level from the source code and the flow description, not an exact value.', 'OVERVIEW')
table(s, ['Question', 'Level 1', 'Level 2', 'Level 3', 'Level 4', 'Level 5'], [
    ['Area (gate equivalents)', 'below 100', '100 to 1,000', '1,000 to 10,000', '10,000 to 100,000', '100,000 or more'],
    ['Power (microwatts)', 'below 10', '10 to 100', '100 to 1,000', '1,000 to 10,000', '10,000 or more'],
    ['Timing margin, slack (ns)', 'negative', '0 to 0.5', '0.5 to 1.0', '1.0 to 1.5', '1.5 or more'],
    ['Reachable frequency (GHz)', 'below 1', '1 to 2', '2 to 3', '3 to 4', '4 or more'],
], .89, 1.98, [3.1, 1.15, 1.4, 1.85, 2.15, 1.9], 2.4, size=14, head_size=14, head_h=.45, row_header=True, nowrap=True)
card(s, .89, 4.6, 5.65, 1.5, 'Units', '1 gate equivalent (GE) = the area of one NAND2 gate = 0.196608 µm² in NanGate 15 nm.\nPower = dynamic + leakage power for the testbench stimulus.', body_size=13.5)
card(s, 6.79, 4.6, 5.65, 1.5, 'Reachable frequency', 'Reachable frequency = 1 / (2 ns − worst slack), from the netlist synthesized at 500 MHz.\nExample: 1.438 ns of slack gives 1 / 0.562 ns = 1.78 GHz.', body_size=13.5)
assert '1' not in DM[M[0]]['table']['slack']['truth_counts'] and '5' not in DM[M[0]]['table']['reach']['truth_counts'] and '5' not in DM[M[0]]['table']['power']['truth_counts']
takeaway(s, 'One level of area or power is a factor of 10. No test design misses the 500 MHz target (slack level 1), reaches 4 GHz, or uses 10 mW or more.')
notes(s, 'The reference synthesis uses RTL_flow_Nangate15: NanGate 15 nm, typical corner at 0.8 V and 25 C, Design Compiler compile_ultra -gate_clock, a 500 MHz target, input and output delay of 10% of the period, and 50 ps setup clock uncertainty. '
      'Area is the total cell area divided by 0.196608 µm², the area of NAND2_X1. Power is the PrimeTime total of dynamic and leakage power with the supplied power testbench; a power value counts only if the switching annotation is at least 95%. '
      'In the test split 85 of the 92 performance requests have a power label; the other 7 are excluded, not model errors. '
      'Worst slack is the smallest timing margin after synthesis at 500 MHz; a positive slack means that the netlist meets the target. Reachable frequency is the harness formula 1000 / (period in ns − slack in ns) MHz of that same netlist; it is not a new synthesis at a faster target. '
      'The example mac_int8 has 1437.97 ps of slack: 1 / (2 − 1.43797) ns = 1.779 GHz. The level ranges are lower-inclusive and upper-exclusive, except the open first and last ranges. '
      f'Test labels per level 1 to 5: area {", ".join(map(str, level_counts("kev-27b", "area", D, "truth_counts")))}; power {", ".join(map(str, level_counts("kev-27b", "power", D, "truth_counts")))}; '
      f'frequency {", ".join(map(str, level_counts("kev-27b", "reach", D, "truth_counts")))}; slack {", ".join(map(str, level_counts("kev-27b", "slack", D, "truth_counts")))}. '
      'So no test design has a negative slack (level 1), a frequency of 4 GHz or more (level 5), or a power of 10 mW or more (level 5).',
      'harness/build.py: AREA_EDGES_GE, POWER_EDGES_UW, SLACK_EDGES_NS, REACH_EDGES_MHZ; notes/DATASET.md; flow/RTL_flow_Nangate15/summary.sh')

# ----------------------------------------------------------------------------------------------
# 6. Experiment 2: the three examples
# ----------------------------------------------------------------------------------------------
VALUE_PATTERN = {'slack': r'worst slack ([+-][\d.]+ ns)', 'reach': r'frequency ([\d,]+ MHz)',
                 'area': r'= ([\d,.]+ GE)', 'power': r'total power ([\d.,]+ [um]W)'}


def example_values(key):
    return '  •  '.join(f"{re.search(VALUE_PATTERN[key], e['result']).group(1).replace('uW', 'µW')} (level {e['level']})"
                        for e in EX[key]['examples'])


def example_answers(key, word, tool):
    a = [e['answer'] for e in EX[key]['examples']]
    return f'{a.count("yes")} {word}, {a.count("no")} fail; each failure with its first {tool} error'


for pair_key, metric_key in [('smaller_area', 'area'), ('lower_power', 'power'), ('higher_reach', 'reach')]:
    assert [e['design'] for e in EX[pair_key]['examples']] == [e['design'] for e in EX[metric_key]['examples']]
s = slide('Experiment 2: three solved examples per question', 'The same 725 test questions, one per request. The three examples are development designs with their tool result.', 'OVERVIEW')
t = table(s, ['Question', 'The three examples before the design, with their tool result'], [
    ['Does it compile?', example_answers('compiles', 'compiles', 'vlog')],
    ['Does it synthesize?', example_answers('synthesizes', 'synthesizes', 'DC')],
    ['Why does synthesis fail?', f'{", ".join(REASON_EX[:-1])} and {REASON_EX[-1]}; each with its first DC error'],
    ['Area level', example_values('area')],
    ['Power level', example_values('power')],
    ['Frequency level', example_values('reach')],
    ['Slack level', example_values('slack')],
    ['Which of two designs?', 'The three designs of the metric (area, power, or frequency), with the exact value only'],
], .89, 1.98, [2.75, 8.8], 4.05, size=15, head_size=15, row_header=True, head_h=.45, left_cols=(0, 1))
takeaway(s, 'All questions of one kind get the same three examples. No example shares a base design with a test design.')
V1_PER_REQUEST = '; '.join(f'{v} {k.split(":")[0]} requests with {k.split(":")[1]} questions' for k, v in sorted(D2['v1_questions_per_test_request'].items())).replace('with 1 questions', 'with 1 question')
notes(s, 'Experiment 2 (MV, 2026-10-16) asks the 725 test questions of experiment 1 again. Each request holds one question; experiment 1 sent up to four questions in one request '
      f'(test split: {V1_PER_REQUEST}). The state of each request is the state of experiment 1 with one new first field, examples; the analysis checks this for all 725 requests. '
      'Kev answers each question in its own row (the state followed by that question), and the Kev README reports that asking questions together or separately gives probabilities within 4e-6 in its fp32 tests. One question per request therefore should not change an answer by itself; the examples are the change. '
      'Rules of the example sets (harness/fewshot.py): compile and synth get the two shortest development designs that the tool rejects (with different errors or classes) and the working base design of the first one; reason gets the three shortest failed designs of three classes; '
      'each level question gets the shortest development design in each level range 1 to 2, 3, and 4 to 5, with the exact value and the level text. A pair question gets the three designs of its metric with the exact value only, no level and no solved pair (MV’s rule). A failed example also shows the first vlog or DC error (an addition that MV accepted). '
      f'No example shares a base design with a test design (checked here); harness/fewshot.py also excludes an example that implements the function of a test design. Exact Kev input tokens: median {D2["tokens"]["median"]:,}, maximum {D2["tokens"]["max"]:,}. '
      f'{len(D2["tokens"]["over_8192"])} requests exceed 8,192 tokens, the validated context of Kev-0.8B, 4B, and 9B (Kev README; Kev-27B: 65,536): {", ".join(q.split(":", 2)[2] for q in D2["tokens"]["over_8192"])}. They are scored like all others.',
      'data/questions/dataset_v1_fewshot3.jsonl; data/questions/dataset_v1_fewshot3.card.md; harness/fewshot.py; notes/VERIFICATION.md item 14; condor/kev_src/README.md')

# ----------------------------------------------------------------------------------------------
# 7. How to read the result charts
# ----------------------------------------------------------------------------------------------
s = slide('How to read the result charts', 'The same chart form on every result slide. The example: the synthesis question.', 'OVERVIEW', legend=True)
yes_no_panel(s, .89, 2.0, 5.65, 4.1, 'synthesizes', 'Does it synthesize?')
card(s, 6.79, 2.0, 5.65, 1.25, 'Two bars per model', 'Light bar: the score without examples (experiment 1).\nDark bar: the score with three solved examples (experiment 2).', body_size=13.5)
card(s, 6.79, 3.42, 5.65, 1.25, 'The label beside a bar', 'The score, then the correct answers on the designs where the tool says “yes”, and on the designs where it says “no”.', body_size=13.5)
card(s, 6.79, 4.84, 5.65, 1.26, 'The dashed lines', 'The score of a model that always answers “no”, and of one that always answers “yes”. A bar at a dashed line shows no skill.', body_size=13.5)
takeaway(s, 'The score is the share of the test questions with a correct answer. Each question has the same weight.')
notes(s, 'Every result slide uses this chart form. One panel is one question. Each model has two bars: the light bar is the score without examples (experiment 1, 15 October 2026) and the dark bar the score with three solved examples (experiment 2, 16 October 2026). '
      'The score is plain accuracy: the share of the test questions of that type with a correct answer (MV, 2026-10-19). For a yes/no question the label gives the score and then the correct answers on the designs where the tool says yes and on the designs where it says no; the two counts add up to the correct answers of the bar. '
      'A yes probability of 0.5 or more counts as yes, exactly as in the saved scorer. The dashed lines are the scores of a constant answer, computed from the test labels: for synthesis always yes scores 80.5% (99 of 123) and always no 19.5% (24 of 123). '
      'For a level question the label gives the score and the count of exact levels, and the dashed line is the most common level of the test labels. For a pair question the dashed line is 50%: both orders of each pair are asked, so a model that always answers A or always B is right exactly once per pair. '
      'A constant answer is a descriptive reference from the test labels, not a separate run and not a predictor chosen on the development split.')

# ----------------------------------------------------------------------------------------------
# 8. The findings in one table
# ----------------------------------------------------------------------------------------------
def score_cell(key, d):
    t = d['models']['kev-27b']['table'][key]
    return f'{pct(t["accuracy"])} ({round(t["accuracy"] * t["n"])} of {t["n"]})'


assert metric('area', 'within_one')[3] == 1.0 and metric('area', 'within_one', D2)[3] == 1.0
assert H[('reason', 1)][3] == 15 and H[('reason', 2)][3] == 14
assert all(H[(k, 2)][3] == max(H[(k, 2)]) and H[(k, 2)][3] > H[(k, 1)][3] for k in ['smaller_area', 'lower_power', 'higher_reach'])
FINDING = {'compiles': 'Reliable without examples', 'synthesizes': 'Close to always “yes”', 'reason': 'Good, when told that it fails',
           'area': 'Usable; all within one level', 'power': 'Close to the constant answer', 'reach': 'Below the constant answer',
           'slack': 'Not above the constant answer', 'smaller_area': 'Good with examples', 'lower_power': 'Good with examples', 'higher_reach': 'Good; only 6 pairs'}
QLABEL = {'compiles': 'Does it compile?', 'synthesizes': 'Does it synthesize?', 'reason': 'Why does synthesis fail? (told that it fails)',
          'area': 'Area level', 'power': 'Power level', 'reach': 'Frequency level', 'slack': 'Slack level',
          'smaller_area': 'Smaller area of two designs', 'lower_power': 'Lower power of two designs', 'higher_reach': 'Higher frequency of two designs'}
CONST = {'compiles': f'“yes”: {pct(N_PASS["compiles"] / 123)}', 'synthesizes': f'“yes”: {pct(N_PASS["synthesizes"] / 123)}',
         'reason': f'one reason: {pct(REASON_MAX / 16)}', 'smaller_area': 'A: 50.0%', 'lower_power': 'A: 50.0%', 'higher_reach': 'A: 50.0%'}
CONST.update({k: f'level {LEVEL_REF[k][0]}: {pct(LEVEL_REF[k][1] / LEVEL_REF[k][2])}' for k in LEVEL_REF})
s = slide('The findings in one table', 'KEV-27B, the largest size. Each score: correct answers in % (count). The other sizes are on the result slides.', 'OVERVIEW')
table(s, ['Question', LEG1.capitalize(), LEG2.capitalize(), 'Always the same answer', 'Finding'],
      [[QLABEL[k], score_cell(k, D), score_cell(k, D2), CONST[k], FINDING[k]] for k in QKEYS],
      .89, 1.98, [3.2, 1.85, 1.85, 1.85, 2.8], 4.15, size=12.5, head_size=13, head_h=.45, left_cols=(0, 4))
takeaway(s, 'KEV-27B is reliable on the compile check, and good on the failure reason and on the better of two designs. It is weak on the synthesis check and on the power and timing levels.')
notes(s, 'The scores of KEV-27B, the largest size, on the ten question keys, without examples and with three examples, with the score of a constant answer. The constant answer is the most common test label of each question: '
      + '; '.join(f'{QLABEL[k]}: {CONST[k]}' for k in QKEYS) + '. '
      f'Compile: {H[("compiles", 1)][3]} of 123 without examples, {H[("compiles", 2)][3]} with examples; the examples add {C27_2["pass_wrong"]} wrong “no” answers on designs that compile. '
      f'Synthesis: {H[("synthesizes", 1)][3]} and {H[("synthesizes", 2)][3]} of 123, against 99 for always yes; the best score of any size and experiment is {BEST["synthesizes"] - 99} designs above always yes. '
      f'Reason: {H[("reason", 1)][3]} and {H[("reason", 2)][3]} of 16, when the prompt says that synthesis failed. '
      f'Area: {H[("area", 1)][3]} and {H[("area", 2)][3]} of 92 exact, all 92 within one level in both experiments. '
      f'Power, frequency, slack: the best exact count of any size and experiment is {BEST["power"]} of 85 (constant {LEVEL_REF["power"][1]}), {BEST["reach"]} of 92 (constant {LEVEL_REF["reach"][1]}), and {BEST["slack"]} of 92 (constant {LEVEL_REF["slack"][1]}; this count is KEV-0.8B with level 5 for every design). '
      f'Pairs with examples: {H[("smaller_area", 2)][3]} of 46, {H[("lower_power", 2)][3]} of 44, {H[("higher_reach", 2)][3]} of 12; KEV-27B has the highest count of the four sizes on each pair question with examples. '
      'A smaller size leads in two places without examples: KEV-9B on the area pairs (40 of 46) and KEV-4B on the exact power level (39 of 85).')

# ----------------------------------------------------------------------------------------------
# 9. The yes/no checks
# ----------------------------------------------------------------------------------------------
GAIN = {(m, e): BIN[('synthesizes', e)][i]['pass_correct'] + BIN[('synthesizes', e)][i]['fail_correct'] - N_PASS['synthesizes'] for i, m in enumerate(M) for e in (1, 2)}
assert max(GAIN.values()) == GAIN[('kev-27b', 2)] == 9 and YES2['synthesizes'][0] == 0 and YES2['compiles'][0] == 1
assert C27_1['pass_correct'] == 114 and C27_1['fail_correct'] == 8 and C27_2['fail_correct'] == 9 and C27_2['pass_wrong'] == 6
s = slide('Yes/no checks: only the compile check is reliable', 'Compile: KEV-27B gets 122 of 123 right. Synthesis: all scores near always “yes”.', 'RESULTS', legend=True)
yes_no_panel(s, .89, 2.0, 5.65, 4.1, 'compiles', 'Does it compile?')
yes_no_panel(s, 6.79, 2.0, 5.65, 4.1, 'synthesizes', 'Does it synthesize?')
takeaway(s, f'Synthesis: the best score, KEV-27B with examples, is {GAIN[("kev-27b", 2)]} designs above always “yes”. KEV-0.8B with examples answers “no” to every design on synthesis.')
notes(s, 'One test set of 123 designs answers both questions. Each bar is the share of the 123 designs with a correct answer; beside it the correct answers on the designs where the tool says yes and where it says no. '
      f'“Yes” answers of 123 without examples: compile {", ".join(map(str, YES1["compiles"]))}, synthesis {", ".join(map(str, YES1["synthesizes"]))}; with examples: compile {", ".join(map(str, YES2["compiles"]))}, synthesis {", ".join(map(str, YES2["synthesizes"]))} (0.8B, 4B, 9B, 27B). '
      f'Compile, KEV-27B: {C27_1["pass_correct"]} of 114 yes and {C27_1["fail_correct"]} of 9 no without examples; with examples {C27_2["pass_correct"]} of 114 and {C27_2["fail_correct"]} of 9, so the examples catch the last compile error but add {C27_2["pass_wrong"]} wrong no answers '
      f'({pval(CMP["kev-27b"]["keys"]["compiles"]["by_class"]["True"]["mcnemar_p"])} within the 114 designs that compile). '
      f'KEV-0.8B answers no to compile for {123 - YES1["compiles"][0]} of 123 designs without examples and {123 - YES2["compiles"][0]} with examples: its 9 of 9 on the compile errors is a near-constant no, not a detection (MV asked about this on 2026-10-16). '
      f'Synthesis, KEV-27B: {S27_1["pass_correct"]} of 99 yes and {S27_1["fail_correct"]} of 24 no without examples; {S27_2["pass_correct"]} of 99 and {S27_2["fail_correct"]} of 24 with examples '
      f'(within the 24 failures: {CMP["kev-27b"]["keys"]["synthesizes"]["by_class"]["False"]["only_2"]} gained, {CMP["kev-27b"]["keys"]["synthesizes"]["by_class"]["False"]["only_1"]} lost, {pval(CMP["kev-27b"]["keys"]["synthesizes"]["by_class"]["False"]["mcnemar_p"])}; all 123: {pval(CMP["kev-27b"]["keys"]["synthesizes"]["mcnemar_p"])}). '
      f'Designs above the nearer constant answer, synthesis: without examples {", ".join(str(GAIN[(m, 1)]) for m in M)}; with examples {", ".join(str(GAIN[(m, 2)]) for m in M)} (0.8B with examples: 0 yes answers, so it sits on the always-no line). '
      'In this test set every design that vlog rejects also fails synthesis. One design changes a score by 0.8 percentage points. The appendix gives every yes probability and the class counts.')

# ----------------------------------------------------------------------------------------------
# 10. Which synthesis failures are caught
# ----------------------------------------------------------------------------------------------
def sub_labels(k, e):
    return [f'{pct(x["correct"] / x["n"])} · {x["correct"]} of {x["n"]}' for x in SUB[(k, e)]]


A27 = (SUB[('compiler_accepts', 1)][3]['correct'], SUB[('compiler_accepts', 2)][3]['correct'])
R27 = (SUB[('compiler_rejects', 1)][3]['correct'], SUB[('compiler_rejects', 2)][3]['correct'])
assert A27 == (0, 4) and R27 == (4, 8) and SUB[('compiler_accepts', 2)][0]['correct'] == 15 and SUB[('compiler_rejects', 2)][0]['correct'] == 9
s = slide('KEV misses most synthesis failures that compile', 'The 24 synthesis failures split by the compile check. Bars: failures answered “no”.', 'RESULTS', legend=True)
gpanel(s, .89, 2.0, 5.65, 4.1, 'Failures whose source compiles', f'{SUB[("compiler_accepts", 1)][0]["n"]} designs: only synthesis rejects them',
       ([x['correct'] / x['n'] for x in SUB[('compiler_accepts', 1)]], [x['correct'] / x['n'] for x in SUB[('compiler_accepts', 2)]]),
       (sub_labels('compiler_accepts', 1), sub_labels('compiler_accepts', 2)))
gpanel(s, 6.79, 2.0, 5.65, 4.1, 'Failures whose source does not compile', f'{SUB[("compiler_rejects", 1)][0]["n"]} designs: vlog rejects them too',
       ([x['correct'] / x['n'] for x in SUB[('compiler_rejects', 1)]], [x['correct'] / x['n'] for x in SUB[('compiler_rejects', 2)]]),
       (sub_labels('compiler_rejects', 1), sub_labels('compiler_rejects', 2)))
takeaway(s, f'KEV-27B catches {A27[1]} of the 15 failures that compile with examples, {A27[0]} without; and {R27[1]} of the 9 that do not compile ({R27[0]} without). KEV-0.8B’s full bars are a constant “no”.')
notes(s, 'A derived subgroup analysis of the synthesis answers: the 24 failing test designs are split by the compile label of the same design. All bars are synthesis yes/no decisions at the 0.5 threshold; the right panel is not the separate compile answer. '
      f'Failures that compile (15), answered no: without examples {", ".join(str(x["correct"]) for x in SUB[("compiler_accepts", 1)])}; with examples {", ".join(str(x["correct"]) for x in SUB[("compiler_accepts", 2)])} (0.8B, 4B, 9B, 27B). '
      f'Failures that do not compile (9): without examples {", ".join(str(x["correct"]) for x in SUB[("compiler_rejects", 1)])}; with examples {", ".join(str(x["correct"]) for x in SUB[("compiler_rejects", 2)])}. '
      'KEV-0.8B with examples gives every one of the 123 designs a synthesis probability below 0.5, so its 15 of 15 and 9 of 9 are a constant no, not a detection. '
      f'Ranking (AUROC of the 99 successes against each subgroup), KEV-27B: {SUB[("compiler_accepts", 2)][3]["auroc_vs_successes"]:.3f} for the 15 and {SUB[("compiler_rejects", 2)][3]["auroc_vs_successes"]:.3f} for the 9 with examples '
      f'(without examples {SUB[("compiler_accepts", 1)][3]["auroc_vs_successes"]:.3f} and {SUB[("compiler_rejects", 1)][3]["auroc_vs_successes"]:.3f}). '
      'The 15 failures that compile are 9 mutants with one injected synthesis error and 6 rejected trap twins; the 9 that do not compile are 8 noncompile edits and 1 natural failure. Small subgroups: no general claim.')

# ----------------------------------------------------------------------------------------------
# 11. The ten failure reasons
# ----------------------------------------------------------------------------------------------
RO, RC = D['reason_options'], D['reason_counts_by_split']
NT, ND = sum(RC['test'].values()), sum(RC['dev'].values())
NR = D['synthesis_failures_without_reason']
assert NR['splits'] == {'test': NR['n']} and NR['classes'] == {'illegal_code': NR['n']} == {'illegal_code': NR['compile_fails']}
s = slide('The ten failure reasons and how often each occurs', f'The prompt says that synthesis failed. {NT + ND} reason questions: {NT} in the test split, {ND} in the development split (dev).', 'RESULTS')
t = table(s, ['Reason', 'Text of the option, as the model reads it', 'Test', 'Dev'],
          [[k, v, RC['test'].get(k, 0), RC['dev'].get(k, 0)] for k, v in RO.items()],
          .89, 1.95, [2.3, 7.65, .7, .9], 4.0, size=10.5, head_size=13, row_header=True, head_h=.4, left_cols=(0, 1))
for i, v in enumerate(RO.values(), 1):
    t.rows[i].height = Inches(.25 if len(v) <= 100 else .40)
takeaway(s, f'{D["reason_injected_edits"]} of the {NT + ND} failures are injected edits. {NR["n"]} more test designs with an injected compile error get no reason question.')
notes(s, f'The ten options of the failure-reason question exactly as the model reads them (option name, then its text). Every one of the {NT + ND} reason questions in dataset v1 offers the same ten options in the same order. '
      f'The columns count the reason questions whose correct answer is that option: {NT} in the test split and {ND} in the development split. In the test split, {len(RC["test"])} of the ten options occur; event_list and unbounded_loop occur only in the development split, once each. '
      'The correct answer comes from the Design Compiler message IDs (and the vlog result to separate illegal_code from unsupported_construct), with the rules in notes/synth_reason.md; a design is kept only if all its messages belong to one class. '
      f'{D["reason_injected_edits"]} of the {NT + ND} failing designs are injected one-construct edits of a working design (mutants and the rejected twins of the trap pairs). The other {NT + ND - D["reason_injected_edits"]} are natural failures: one MetRex design and two LLM-written candidates. '
      f'{NR["n"]} further test designs fail synthesis with illegal_code; they are the injected compile errors (one removed semicolon or one undeclared name) and get only the yes/no questions, by the dataset rules. MV asked for this table on 2026-10-16.',
      'data/questions/dataset_v1.jsonl (records with request = reason); notes/synth_reason.md; notes/DATASET.md; presentations/data/verified_metrics.json')

# ----------------------------------------------------------------------------------------------
# 12. The failure reason: naming a known failure against detecting it
# ----------------------------------------------------------------------------------------------
assert RL[1][3]['detected'] == 0 and RL[2][3]['detected'] == 5 and RL[1][3]['reason_correct'] == 15 and RL[2][3]['reason_correct'] == 14
assert RL[2][0]['detected'] == 16 and YES2['synthesizes'][0] == 0
R27 = D2M['kev-27b']['reason_by_class']
WRONG27 = {k: sorted(p for p in v['predictions'] if p != k) for k, v in R27.items() if v['correct'] < v['n']}
assert WRONG27 == {'multiple_drivers': ['event_list'], 'async_reset': ['event_list']}
EVENT2 = [sum(v['predictions'].get('event_list', 0) for v in D2M[m]['reason_by_class'].values()) for m in M]
EVENT1 = [CMP[m]['reason_predictions_1'].get('event_list', 0) for m in M]
s = slide('KEV-27B names a known failure but does not detect it', 'Left: the reason question. Right: synthesis and reason answers of the 16 designs.', 'RESULTS', legend=True)
gpanel(s, .89, 2.0, 5.65, 4.1, 'Why does synthesis fail?', '16 designs · the prompt says that synthesis failed',
       (metric('reason'), metric('reason', d=D2)), (count_labels('reason', D), count_labels('reason', D2)),
       reference=(REASON_MAX / 16, f'always {REASON_MODE[0]} = {pct(REASON_MAX / 16)} ({REASON_MAX} of 16)'))
table(s, ['Model', 'Says that\nsynthesis fails', 'Names the\ncorrect reason', 'Both'],
      [[NAMES[i], change(RL[1][i]['detected'], RL[2][i]['detected']), change(RL[1][i]['reason_correct'], RL[2][i]['reason_correct']), change(RL[1][i]['both'], RL[2][i]['both'])] for i in range(4)],
      6.79, 2.0, [1.45, 1.5, 1.5, 1.2], 2.7, size=15, head_size=13, head_h=.62)
txt(s, 'Each cell: designs of 16, no examples → three examples.\nThe synthesis answer comes from the yes/no request of the same design, where the prompt does not say that it fails.',
    6.85, 4.85, 5.6, 1.2, size=13, color=C['muted'])
takeaway(s, f'KEV-27B names the reason for {RL[1][3]["reason_correct"]} of 16 known failures without examples, but it says “synthesis fails” for {RL[1][3]["detected"]} of these 16 designs when the prompt does not tell it ({RL[2][3]["detected"]} with examples).')
notes(s, f'Left panel: correct reason choices of 16 without examples {", ".join(map(str, H[("reason", 1)]))}, with examples {", ".join(map(str, H[("reason", 2)]))} (0.8B, 4B, 9B, 27B). The dashed line is the most common class, {REASON_MODE[0]} ({REASON_MAX} of 16). '
      'The instruction states the failure, so this score is conditional on a known failure; it is not detection. KEV-27B without examples is wrong on the one multiple_drivers design (it chooses async_reset); with examples it is wrong on the multiple_drivers design and the async_reset design, and chooses event_list for both. '
      f'event_list is one of the three example classes and the class of no test design; the models choose it {sum(EVENT2)} times with examples ({", ".join(map(str, EVENT2))}) and {sum(EVENT1)} times without ({", ".join(map(str, EVENT1))}). This is an observation on one example set; the cause is not tested. '
      'Right table: for each of the 16 reason-test designs, the synthesis yes/no answer of the same design (from its own request, where the prompt does not state the failure) is matched with the reason answer. '
      f'Says that synthesis fails: {", ".join(str(r["detected"]) for r in RL[1])} without examples, {", ".join(str(r["detected"]) for r in RL[2])} with examples; names the reason: {", ".join(str(r["reason_correct"]) for r in RL[1])} and {", ".join(str(r["reason_correct"]) for r in RL[2])}; both: {", ".join(str(r["both"]) for r in RL[1])} and {", ".join(str(r["both"]) for r in RL[2])}. '
      'KEV-0.8B with examples answers no to the synthesis question for every one of the 123 designs, so its 16 of 16 carries no information. The reason subset has 16 known failures; the full synthesis test has 24 failures, because 8 noncompile designs get no reason request. '
      'These are matched saved answers, not a deployed detection-and-diagnosis pipeline.')

# ----------------------------------------------------------------------------------------------
# 13. Area and power levels
# ----------------------------------------------------------------------------------------------
AW = (round(metric('area', 'within_one')[3] * 92), round(metric('area', 'within_one', D2)[3] * 92))
assert AW == (92, 92) and H[('area', 1)][3] == 61 and H[('area', 2)][3] == 48 and one_answer('kev-0.8b', 'area', D2) == '5' and one_answer('kev-0.8b', 'power', D2) == '5'
assert CMP['kev-27b']['keys']['area']['mcnemar_p'] < .05 and CMP['kev-27b']['keys']['power']['mcnemar_p'] < .05
s = slide('Area levels are usable; power levels are not', 'Bars: exact level. Dashed line: the most common level of the tool labels.', 'RESULTS', legend=True)
level_panel(s, .89, 2.0, 5.65, 4.1, 'area', 'Area level', '92 designs · all five levels occur')
level_panel(s, 6.79, 2.0, 5.65, 4.1, 'power', 'Power level', '85 designs · levels 1 to 4 occur')
takeaway(s, f'Area: KEV-27B is exact for {H[("area", 1)][3]} of 92 designs and within one level for all 92; with examples the exact count drops to {H[("area", 2)][3]}. Power: only KEV-4B without examples ({H[("power", 1)][1]} of 85) is above always level {LEVEL_REF["power"][0]} ({LEVEL_REF["power"][1]} of 85).')
W1 = {k: [round(v * n_of(k)) for v in metric(k, 'within_one')] for k in ['area', 'power']}
W2 = {k: [round(v * n_of(k)) for v in metric(k, 'within_one', D2)] for k in ['area', 'power']}
notes(s, f'Exact area levels of 92: without examples {", ".join(map(str, H[("area", 1)]))}, with examples {", ".join(map(str, H[("area", 2)]))} (0.8B, 4B, 9B, 27B); within one level: {", ".join(map(str, W1["area"]))} and {", ".join(map(str, W2["area"]))}. '
      f'KEV-27B: {H[("area", 1)][3]} to {H[("area", 2)][3]} exact ({pval(CMP["kev-27b"]["keys"]["area"]["mcnemar_p"])}, paired McNemar test), all 92 within one level in both experiments. The dashed line is always level 1, the most common test label ({LEVEL_REF["area"][1]} of 92). '
      f'Exact power levels of 85: without examples {", ".join(map(str, H[("power", 1)]))}, with examples {", ".join(map(str, H[("power", 2)]))}; within one level: {", ".join(map(str, W1["power"]))} and {", ".join(map(str, W2["power"]))}. '
      f'KEV-27B: {H[("power", 1)][3]} to {H[("power", 2)][3]} exact ({pval(CMP["kev-27b"]["keys"]["power"]["mcnemar_p"])}); KEV-4B has the best count without examples ({H[("power", 1)][1]}). Always level 2 scores {LEVEL_REF["power"][1]} of 85; only the KEV-4B bar without examples is above it, by 2 designs. '
      'KEV-0.8B with examples chooses level 5 for all 92 area designs and all 85 power designs, so its exact scores are 2 of 92 and 0 of 85. '
      'The level ranges are decades, so a within-one answer can still be off by a factor of 10. Power depends on the synthesized design and on the stimulus of its testbench, which the model receives. '
      f'The examples of the area question: {example_values("area")}; of the power question: {example_values("power")}.')

# ----------------------------------------------------------------------------------------------
# 14. Timing levels
# ----------------------------------------------------------------------------------------------
assert H[('reach', 2)][3] == 36 and H[('slack', 2)][3] == 60 and H[('slack', 2)][0] == LEVEL_REF['slack'][1] == 63 and one_answer('kev-0.8b', 'slack', D2) == '5'
s = slide('Timing levels: no model beats a constant answer', 'Bars: exact level. Dashed line: the most common tool level. No slack is negative.', 'RESULTS', legend=True)
level_panel(s, .89, 2.0, 5.65, 4.1, 'reach', 'Reachable clock frequency level', '92 designs · levels 1 to 4 occur')
level_panel(s, 6.79, 2.0, 5.65, 4.1, 'slack', 'Timing margin (slack) level', '92 designs · levels 2 to 5 occur')
takeaway(s, f'With examples KEV-27B reaches {H[("reach", 2)][3]} of 92 (frequency) and {H[("slack", 2)][3]} of 92 (slack); the constant answers score {LEVEL_REF["reach"][1]} and {LEVEL_REF["slack"][1]}. KEV-0.8B’s {H[("slack", 2)][0]} of 92 is level 5 for every design.')
notes(s, f'Exact frequency levels of 92: without examples {", ".join(map(str, H[("reach", 1)]))}, with examples {", ".join(map(str, H[("reach", 2)]))} (0.8B, 4B, 9B, 27B). Exact slack levels of 92: {", ".join(map(str, H[("slack", 1)]))} and {", ".join(map(str, H[("slack", 2)]))}. '
      f'The constant answers are the most common test labels: frequency level {LEVEL_REF["reach"][0]} ({LEVEL_REF["reach"][1]} of 92) and slack level {LEVEL_REF["slack"][0]} ({LEVEL_REF["slack"][1]} of 92). '
      f'KEV-27B improves with examples on both (frequency {pval(CMP["kev-27b"]["keys"]["reach"]["mcnemar_p"])}, slack {pval(CMP["kev-27b"]["keys"]["slack"]["mcnemar_p"])}) and stays below the constant answer. '
      'KEV-0.8B with examples chooses slack level 5 for all 92 designs; it ties the constant answer and does not beat it. '
      'Both timing levels come from the same netlist synthesized at 500 MHz, so the two questions are linked: the reachable frequency is 1 / (2 ns − slack). No test design misses the 500 MHz target, so missed timing violations are not tested. '
      f'The examples of the frequency question: {example_values("reach")}; of the slack question: {example_values("slack")}. Level accuracy is not a timing error in picoseconds or MHz.')

# ----------------------------------------------------------------------------------------------
# 15. The chosen levels of KEV-27B
# ----------------------------------------------------------------------------------------------
T27 = {k: (level_counts('kev-27b', k, D, 'truth_counts'), level_counts('kev-27b', k, D), level_counts('kev-27b', k, D2)) for k in ['reach', 'power']}
assert T27['reach'][:2] == ([13, 16, 18, 45, 0], [90, 0, 0, 0, 2]) and T27['power'][:2] == ([23, 37, 18, 7, 0], [0, 0, 54, 30, 1])


def gap(a, b):
    return sum(abs(x - y) for x, y in zip(a, b))


assert all(gap(T27[k][2], T27[k][0]) < gap(T27[k][1], T27[k][0]) for k in T27)
s = slide('KEV-27B’s chosen levels do not match the tool’s levels', 'Designs at each level: the tool, then KEV-27B without and with examples. Equal counts are not correct answers.', 'RESULTS')
dist3(s, .89, 2.0, 5.65, 4.1, 'Reachable clock frequency level', '92 designs · counts at each level', *T27['reach'])
dist3(s, 6.79, 2.0, 5.65, 4.1, 'Power level', '85 designs · counts at each level', *T27['power'])
takeaway(s, f'Without examples KEV-27B puts {T27["reach"][1][0]} of 92 designs in frequency level 1 and no design in power level 1 or 2. With examples the counts come closer to the tool’s, but the exact scores stay below the constant answer.')
notes(s, f'KEV-27B, frequency: tool counts at levels 1 to 5 {", ".join(map(str, T27["reach"][0]))}; chosen without examples {", ".join(map(str, T27["reach"][1]))}; with examples {", ".join(map(str, T27["reach"][2]))}. '
      f'Power: tool {", ".join(map(str, T27["power"][0]))}; chosen without examples {", ".join(map(str, T27["power"][1]))}; with examples {", ".join(map(str, T27["power"][2]))}. '
      f'“Closer” means a smaller sum over the five levels of |chosen count − tool count|: frequency {gap(T27["reach"][2], T27["reach"][0])} with examples against {gap(T27["reach"][1], T27["reach"][0])} without; power {gap(T27["power"][2], T27["power"][0])} against {gap(T27["power"][1], T27["power"][0])}. '
      'Equal counts do not mean correct answers: a design can get a level of the right count but the wrong design; the exact scores are on the two previous slides. '
      'Without examples KEV-27B chooses the lowest frequency level for nearly every design and only the high power levels 3 to 5, although 60 of the 85 power labels are levels 1 and 2. These counts show a bias in the chosen levels, not its cause. Both charts use the same 0 to 100 scale.')

# ----------------------------------------------------------------------------------------------
# 16. Pairs
# ----------------------------------------------------------------------------------------------
BOTH = {e: {k: PC[e]['kev-27b'][k]['both_correct'] for k in PAIRS} for e in (1, 2)}
assert BOTH[1]['smaller_area'] == 17 and BOTH[2]['smaller_area'] == 20 and H[('smaller_area', 1)][2] == 40 == max(H[('smaller_area', 1)])
LET = {m: D2M[m]['pair_letters'] for m in M}
assert LET['kev-0.8b']['smaller_area'] == {'A': 46} and LET['kev-0.8b']['higher_reach'] == {'B': 12} and LET['kev-4b']['higher_reach'] == {'B': 12}
s = slide('Pairs: KEV-27B with examples picks the better design', 'Two designs of one problem, A and B, asked in both orders. Always A scores 50%.', 'RESULTS', legend=True)
for x, key, title in [(.89, 'smaller_area', 'Smaller area'), (4.81, 'lower_power', 'Lower power'), (8.74, 'higher_reach', 'Higher frequency')]:
    n = n_of(key)
    gpanel(s, x, 2.0, 3.7, 4.1, title, f'{PAIRS[key]} pairs · {n} questions', (metric(key), metric(key, d=D2)),
           (count_labels(key, D), count_labels(key, D2)), reference=(.5, 'always A = 50%'), label_w=1.3)
takeaway(s, f'With examples KEV-27B is right for {H[("smaller_area", 2)][3]} of 46 area, {H[("lower_power", 2)][3]} of 44 power, and {H[("higher_reach", 2)][3]} of 12 frequency questions, and for {BOTH[2]["smaller_area"]} of 23 area pairs in both orders ({BOTH[1]["smaller_area"]} without).')
notes(s, 'Pair inputs are two REvolution candidates of the same problem with the same flow and the same power testbench; a question is kept for a metric if the two values differ by at least 10%, so there are no ties. Each pair is sent as A/B and as B/A, so the correct letters are balanced and a constant letter scores exactly 50%. '
      f'Correct answers without examples: area {", ".join(map(str, H[("smaller_area", 1)]))} of 46, power {", ".join(map(str, H[("lower_power", 1)]))} of 44, frequency {", ".join(map(str, H[("higher_reach", 1)]))} of 12 (0.8B, 4B, 9B, 27B); '
      f'with examples: area {", ".join(map(str, H[("smaller_area", 2)]))}, power {", ".join(map(str, H[("lower_power", 2)]))}, frequency {", ".join(map(str, H[("higher_reach", 2)]))}. '
      f'Pairs right in both orders, KEV-27B: area {BOTH[1]["smaller_area"]} to {BOTH[2]["smaller_area"]} of 23, power {BOTH[1]["lower_power"]} to {BOTH[2]["lower_power"]} of 22, frequency {BOTH[1]["higher_reach"]} to {BOTH[2]["higher_reach"]} of 6 (the appendix has all sizes). '
      f'Paired McNemar tests of KEV-27B: area {pval(CMP["kev-27b"]["keys"]["smaller_area"]["mcnemar_p"])}, power {pval(CMP["kev-27b"]["keys"]["lower_power"]["mcnemar_p"])}, frequency {pval(CMP["kev-27b"]["keys"]["higher_reach"]["mcnemar_p"])}: KEV-27B has the highest count with examples, but no change against experiment 1 has p < 0.05. '
      'Without examples KEV-9B leads the area pairs with 40 of 46. KEV-0.8B with examples answers A for all 46 area questions and B for all 12 frequency questions; KEV-4B answers B for all 12 frequency questions. '
      'The two orders share designs and different pairs can share candidates, so the questions are not independent samples. The pair questions show the three examples of their metric with the exact values only.')

# ----------------------------------------------------------------------------------------------
# 17. Conclusion
# ----------------------------------------------------------------------------------------------
s = slide('What the results support, and what to do next', 'Observations on dataset v1 and one fixed example set, not a guarantee for other designs or other examples.', 'CONCLUSION')
card(s, .89, 2.0, 3.7, 4.1, 'Supported', f'• Compile check: KEV-27B, {H[("compiles", 1)][3]} of 123, no examples needed.\n'
     f'• Failure reason: KEV-27B, {H[("reason", 1)][3]} of 16, when the prompt says that synthesis failed.\n'
     f'• Area level: KEV-27B, {H[("area", 1)][3]} of 92 exact, all 92 within one level.\n'
     f'• Better of two designs: KEV-27B with examples, {H[("smaller_area", 2)][3]} of 46 (area), {H[("lower_power", 2)][3]} of 44 (power).', C['navy'], 14)
card(s, 4.81, 2.0, 3.7, 4.1, 'Not supported', f'• Synthesis check: at most {GAIN[("kev-27b", 2)]} designs above always “yes”; KEV-27B misses {S27_2["fail_wrong"]} of 24 failures with examples, {S27_1["fail_wrong"]} without.\n'
     '• Power, frequency, and slack levels: no score more than 2 designs above the constant answer.\n'
     f'• KEV-0.8B: one answer for every design in {len(CONST_08)} of 10 questions with examples.', C['teal'], 14)
card(s, 8.74, 2.0, 3.7, 4.1, 'Next steps', '• Dataset v2: more synthesis failures of code that compiles, designs that miss the timing target, equal yes and no counts.\n'
     '• Other example sets, and more than three examples.\n'
     '• Jev 1.13 and other models under the same protocol.\n'
     '• Thresholds and example sets chosen on the development split only.', C['blue'], 14)
takeaway(s, f'KEV does not replace the tools: the best synthesis check misses {S27_2["fail_wrong"]} of 24 failures. KEV can screen compile errors and rank two designs.')
notes(s, 'Supported: the measured scores of KEV-27B that are far above the constant answer in at least one experiment: the compile check, the failure reason with a known failure, the exact area level, and the pair questions. '
      f'Not supported: the synthesis check (the best score of any size and experiment is {GAIN[("kev-27b", 2)]} designs above always yes), the power, frequency, and slack levels (only KEV-4B on power without examples is above the dashed line, by 2 designs), and KEV-0.8B as a whole '
      f'(with examples it gives one answer to every design in {len(CONST_08)} of 10 question keys: {", ".join(CONST_08)}). '
      'The yes/no examples hold one yes and two no designs (MV’s setup); the effect of this mix is not measured. event_list is the class of one reason example and of no test design. '
      'Any threshold or example-set choice must use the development split, not these test labels. The test set has no negative-slack design, no level-5 power design, and no level-5 frequency design. '
      'Not measured: Jev, base Qwen, commercial LLMs, renamed or commented controls, energy per operation, post-route timing, and measured silicon.')

# ----------------------------------------------------------------------------------------------
# Appendix
# ----------------------------------------------------------------------------------------------
s = slide('', section='APPENDIX', bg=C['navy'])
for sh in list(s.shapes):
    sh._element.getparent().remove(sh._element)
TEXT_BOXES[:] = [b for b in TEXT_BOXES if b['slide'] != len(prs.slides)]
txt(s, 'APPENDIX', .89, 2.0, 11.2, .4, size=16, color=C['cyan'], bold=True)
txt(s, 'Full results and\nreproduction', .89, 2.7, 10.9, 1.65, size=43, color=C['white'], bold=True)
txt(s, 'Every score of every size, the class counts, the probabilities, the level errors, the pair orders, the runs, and the files', .89, 5.05, 11.2, .75, size=23, color='BFDDEC')
SLIDE_META[-1]['title'] = 'Full results and reproduction'
notes(s, 'Appendix for review and reproduction. The tables give all primary scores of the four sizes in both experiments with the paired test of each change, the yes/no class counts, every yes probability, the reason classes, the probability metrics, the level tolerance and error, the pair order counts, the design groups, the Condor runs, and the files with their checksums. '
      'There is no single overall score across the question types, because the types have different meanings, counts, and label mixes.')

# A2. Every score, both experiments, with the paired test.
LABELS = {'compiles': 'Does it compile?', 'synthesizes': 'Does it synthesize?', 'reason': 'Failure reason (told)', 'area': 'Area: exact level', 'power': 'Power: exact level',
          'reach': 'Frequency: exact level', 'slack': 'Slack: exact level', 'smaller_area': 'Pair: smaller area', 'lower_power': 'Pair: lower power', 'higher_reach': 'Pair: higher frequency'}


def change_cell(m, k):
    c = CMP[m]['keys'][k]
    return f'{pct(c["correct_1"] / c["n"])} → {pct(c["correct_2"] / c["n"])}'


s = slide('Every score, without and with three examples', f'Each cell: correct answers in %, no examples → three examples. Bold: p < 0.05 ({len(SIG)} of {len(M) * len(QKEYS)} changes).', 'APPENDIX')
t = table(s, ['Question', 'n'] + SHORT + ['Constant'], [[LABELS[k], n_of(k)] + [change_cell(m, k) for m in M] + [CONST[k].split(': ')[1]] for k in QKEYS],
          .89, 1.98, [2.6, .65, 1.75, 1.75, 1.75, 1.75, 1.3], 4.3, size=12.5, head_size=13, head_h=.45)
for i, k in enumerate(QKEYS, 1):
    for j, m in enumerate(M, 2):
        for p in t.cell(i, j).text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = CMP[m]['keys'][k]['mcnemar_p'] < .05
txt(s, 'Constant: the score of the most common test label (always “yes”, one level, one reason, or one letter). The p value is the exact two-sided McNemar test on the paired questions; with 40 tests about 2 fall below 0.05 by chance.',
    .95, 6.4, 11.4, .5, size=12, color=C['muted'])
notes(s, 'The complete primary table of both experiments from the saved scorer, reproduced from the raw answers. Yes/no scores use the threshold 0.5. The reason score is conditional on the prompt stating the failure. Area, power, frequency, and slack require the exact level. Pair scores count both orders of each pair as separate questions. '
      'For each size and question key the questions are matched one to one (the question ID of experiment 1 is in the field source_question_id of experiment 2). b = questions correct only without examples, c = questions correct only with examples; '
      'the exact two-sided McNemar p value is the binomial test of b against b + c with probability 0.5 (the rule of harness/compare.py; the analysis recomputes it and checks it against each compare_with_dataset_v1.json). '
      f'p < 0.05: {", ".join(f"{SHORT[M.index(m)]} {LABELS[k]}" for m, k in SIG)}. The KEV-0.8B changes with p < 0.05 come with one answer for every design (synthesis no; area, power, and slack level 5). The test counts changed answers only; it says nothing about the reason for a change. No correction for the 40 tests is applied. '
      'Counts per cell: ' + '; '.join(f'{SHORT[i]} ' + ', '.join(f'{LABELS[k]} {CMP[m]["keys"][k]["correct_1"]} → {CMP[m]["keys"][k]["correct_2"]} of {CMP[m]["keys"][k]["n"]} ({pval(CMP[m]["keys"][k]["mcnemar_p"])})' for k in QKEYS) for i, m in enumerate(M)) + '.')

# A3. Yes/no class counts.
s = slide('Yes/no checks: correct answers per tool outcome', 'Each cell: designs of the group with a correct answer, no examples → three examples. “Yes” answers: of 123.', 'APPENDIX')
rows = []
for key, label in [('compiles', 'Compile'), ('synthesizes', 'Synthesis')]:
    for outcome, word in [('pass', 'the tool says “yes”'), ('fail', 'the tool says “no”')]:
        n = BIN[(key, 1)][0]['n_' + outcome]
        rows.append([f'{label}: {word}', n] + [change(BIN[(key, 1)][i][outcome + '_correct'], BIN[(key, 2)][i][outcome + '_correct']) for i in range(4)])
    rows.append([f'{label}: all designs', 123] + [change(H[(key, 1)][i], H[(key, 2)][i]) for i in range(4)])
    rows.append([f'{label}: “yes” answers', 123] + [change(YES1[key][i], YES2[key][i]) for i in range(4)])
table(s, ['Group', 'Designs'] + SHORT, rows, .89, 1.98, [3.4, 1.15, 1.75, 1.75, 1.75, 1.75], 4.1, size=15, head_size=14, head_h=.45)
txt(s, '“All designs” is the score of the result slides. The paired tests within each group are in the speaker notes.', .95, 6.25, 11.4, .3, size=12.5, color=C['muted'])
notes(s, 'The four binary groups (compile: 114 designs that compile and 9 that do not; synthesis: 99 that synthesize and 24 that do not) with the correct answers of each size in both experiments; wrong counts are the group size minus the correct count. '
      'The all-designs rows are plain accuracy, the score of the result slides. The yes rows count the answers with a yes probability of 0.5 or more among all 123 designs: correct answers on the yes group plus wrong answers on the no group. '
      'Paired McNemar tests within each group, no examples against three examples: ' + '; '.join(
          f'{SHORT[i]} {key} {"yes" if truth == "True" else "no"} group {c["only_2"]} gained, {c["only_1"]} lost ({pval(c["mcnemar_p"])})'
          for i, m in enumerate(M) for key in ['compiles', 'synthesizes'] for truth, c in CMP[m]['keys'][key]['by_class'].items()) + '. '
      'KEV-0.8B answers yes to synthesis for all 123 designs without examples and no for all 123 with examples. KEV-4B answers yes to compile for all 123 designs without examples.')

# A4, A5. Every yes probability.
for e, d, word in [(1, D, 'no examples'), (2, D2, 'three examples')]:
    s = slide(f'Yes probability of every design, {word}', 'Each dot is one test design. 0.5 or more counts as “yes”. Teal: the tool says “yes”; orange: it says “no”.', 'APPENDIX')
    strip_panel(s, .89, 1.99, 5.65, 4.08, 'Does it compile?', 'compiles', ['vlog accepts', 'vlog rejects'], d)
    strip_panel(s, 6.79, 1.99, 5.65, 4.08, 'Does it synthesize?', 'synthesizes', ['tool synthesizes', 'tool fails'], d)
    right_both = [(m, key) for m in M for key in ['compiles', 'synthesizes'] if d['models'][m]['binary'][key]['pass_rate'] > .5 and d['models'][m]['binary'][key]['fail_rate'] > .5]
    assert right_both == [('kev-27b', 'compiles')], right_both
    b27 = d['models']['kev-27b']['binary']['compiles']
    takeaway(s, f'Only KEV-27B’s compile answers are right for most designs of both groups: {b27["pass_correct"]} of {b27["n_pass"]} and {b27["fail_correct"]} of {b27["n_fail"]}.')
    p08 = {k: [p for g in ['pass', 'fail'] for _, p in d['models']['kev-0.8b']['noul'][k][g]] for k in ['compiles', 'synthesizes']}
    sub27 = d['models']['kev-27b']['synthesis_failure_subsets']
    yes = {k: [d['models'][m]['binary'][k]['yes_answers'] for m in M] for k in ['compiles', 'synthesizes']}
    notes(s, f'Each dot is the saved noul answer (the yes probability) of one test design in experiment {e}. In each model row, the upper teal dots are the designs that the tool accepts and the lower orange dots the designs that the tool rejects. '
          'The small vertical spread in each lane comes from a fixed hash of the question ID and only reduces overlap; it carries no data. A probability of 0.5 or more counts as yes, as in the saved scorer. '
          f'“Yes” answers of 123: compile {", ".join(map(str, yes["compiles"]))}, synthesis {", ".join(map(str, yes["synthesizes"]))} (0.8B, 4B, 9B, 27B). '
          f'KEV-0.8B probabilities: compile {min(p08["compiles"]):.2f} to {max(p08["compiles"]):.2f}, synthesis {min(p08["synthesizes"]):.2f} to {max(p08["synthesizes"]):.2f}, for successes and failures alike '
          f'(AUROC {d["models"]["kev-0.8b"]["table"]["compiles"]["auroc"]:.3f} and {d["models"]["kev-0.8b"]["table"]["synthesizes"]["auroc"]:.3f}; 0.5 is chance). '
          f'KEV-27B synthesis: the {sub27["compiler_rejects"]["n"]} failures that also fail vlog: {sub27["compiler_rejects"]["correct"]} below 0.5 (AUROC against the successes {sub27["compiler_rejects"]["auroc_vs_successes"]:.3f}); '
          f'the {sub27["compiler_accepts"]["n"]} failures that pass vlog: {sub27["compiler_accepts"]["correct"]} below 0.5, probabilities {sub27["compiler_accepts"]["min_p"]:.2f} to {sub27["compiler_accepts"]["max_p"]:.2f} (AUROC {sub27["compiler_accepts"]["auroc_vs_successes"]:.3f}). '
          + ('The Kev README names no RTL training data. A search of the Kev suites (evals/) for Verilog terms finds only the C++ logging macro VLOG and the word vlog in a text about videos. Kev-27B also trained on C++ code-review questions (devtools-v1) and starts from a post-trained Qwen model whose training data is not published.' if e == 1 else ''),
          'results/kev/*/dataset_v1*__*.jsonl; presentations/src/analyse_kev.py; github.com/jaredpalmer/kev README and evals/ (commit 5e42a7a)')

# A6. Reason classes.
REASON_NAMES = {'multiple_drivers': 'More than one driver', 'illegal_code': 'Illegal code', 'unsupported_construct': 'Unsupported construct', 'timing_control': 'Unsupported timing control',
                'simulation_statement': 'Simulation-only statement', 'async_reset': 'Asynchronous reset', 'mixed_assignments': 'Mixed assignment styles', 'instance_mismatch': 'Module-instance mismatch'}
s = slide('Failure reasons: correct answers per class', 'The 16 test designs by their true reason. Each cell: correct answers, no examples → three examples.', 'APPENDIX')
rows = [[REASON_NAMES[k], n] + [change(DM[m]['reason_by_class'][k]['correct'], D2M[m]['reason_by_class'][k]['correct']) for m in M] for k, n in D['test_reason_counts'].items()]
rows.append(['Equal weight per class (balanced accuracy)', ''] + [f'{pct(a)} → {pct(b)}' for a, b in zip(metric('reason', 'balanced_accuracy'), metric('reason', 'balanced_accuracy', D2))])
table(s, ['True reason', 'Designs'] + SHORT, rows, .89, 1.98, [3.9, 1.0, 1.6625, 1.6625, 1.6625, 1.6625], 4.1, size=14, head_size=14, head_h=.45)
txt(s, f'No test design has the reason event_list or unbounded_loop. The models chose event_list {sum(EVENT1)} times without examples and {sum(EVENT2)} times with examples ({", ".join(map(str, EVENT2))}); it is one of the three example classes.',
    .95, 6.2, 11.4, .55, size=12.5, color=C['muted'])
WRONG2 = {SHORT[i]: {k: {p: c for p, c in v['predictions'].items() if p != k} for k, v in D2M[m]['reason_by_class'].items() if v['correct'] < v['n']} for i, m in enumerate(M)}
notes(s, 'The eight observed reason classes of the 16 test designs, with the correct answers of each size in both experiments. The scorer’s balanced accuracy weights the eight observed class recalls equally; event_list and unbounded_loop have no test design and are excluded from it. '
      'Five simulation-only designs and three timing-control designs give those classes more weight in the plain score than the one-design classes. '
      'Wrong answers with examples, per size (true class: chosen class and count): ' + '; '.join(f'{k}: ' + ', '.join(f'{t} → ' + ', '.join(f'{p} ×{c}' for p, c in w.items()) for t, w in v.items()) for k, v in WRONG2.items()) + '. '
      f'event_list answers per size: {", ".join(map(str, EVENT1))} without examples, {", ".join(map(str, EVENT2))} with examples. Do not read this as a broad ten-class validation.',
      'notes/synth_reason.md; data/questions/dataset_v1.jsonl; presentations/data/verified_metrics.json; presentations/data/verified_metrics_fewshot3.json')

# A7. Probability metrics.
AUC27 = (DM['kev-27b']['table']['synthesizes']['auroc'], D2M['kev-27b']['table']['synthesizes']['auroc'])
assert AUC27[1] > AUC27[0]
s = slide('Ranking score and probability error, yes/no checks', 'A model can rank well and still decide badly at the 0.5 threshold. Cells: no examples → three examples.', 'APPENDIX')
rows = []
for key, label in [('compiles', 'Compile'), ('synthesizes', 'Synthesis')]:
    rows.append([f'{label}: ranking score (AUROC)'] + [f'{a:.3f} → {b:.3f}' for a, b in zip(metric(key, 'auroc'), metric(key, 'auroc', D2))])
    rows.append([f'{label}: probability error (Brier)'] + [f'{a:.3f} → {b:.3f}' for a, b in zip(metric(key, 'brier'), metric(key, 'brier', D2))])
table(s, ['Probability metric'] + SHORT, rows, .89, 2.0, [3.95, 1.9, 1.9, 1.9, 1.9], 2.35, size=15, head_size=14, head_h=.45)
card(s, .89, 4.65, 5.65, 1.45, 'Ranking score: higher is better', 'Chance = 0.5; perfect ordering = 1. The probability that a design that works gets a higher yes probability than a design that fails.', body_size=13.5)
card(s, 6.79, 4.65, 5.65, 1.45, 'Probability error: lower is better', 'Mean squared error of the yes probability. Zero is perfect; a confident wrong answer costs more than an unsure one.', body_size=13.5)
takeaway(s, f'KEV-27B’s synthesis ranking score rises from {AUC27[0]:.3f} to {AUC27[1]:.3f} with examples, while its failure detection at 0.5 rises from {S27_1["fail_correct"]} to {S27_2["fail_correct"]} of 24.')
notes(s, 'AUROC is the probability that a randomly chosen design that the tool accepts has a higher yes probability than a randomly chosen design that the tool rejects, ties counted as half. It ranges from 0 to 1; 0.5 is chance. It does not depend on the decision threshold and is not failure recall. '
      'The Brier score is the mean of (yes probability − truth)², from 0 to 1, lower is better; it measures probability error, not calibration alone. Both use all 123 test designs. '
      f'KEV-27B synthesis AUROC {AUC27[0]:.3f} to {AUC27[1]:.3f}; its failure recall at 0.5 is {pct(S27_1["fail_rate"])} to {pct(S27_2["fail_rate"])}. Any threshold tuning must use the development split, not these test labels; no recalibrated or alternative-threshold result has been run.')

# A8. Within one level and mean level error.
s = slide('Within one level, and mean level error', 'Within one level: higher is better. Mean level error: lower is better. Each cell: no examples → three examples.', 'APPENDIX')
rows = []
for key, label in [('area', 'Area'), ('power', 'Power'), ('reach', 'Frequency'), ('slack', 'Slack')]:
    rows.append([label, 'Within one level'] + [f'{pct(a)} → {pct(b)}' for a, b in zip(metric(key, 'within_one'), metric(key, 'within_one', D2))])
    rows.append([label, 'Mean level error'] + [f'{a:.2f} → {b:.2f}' for a, b in zip(metric(key, 'mae_levels'), metric(key, 'mae_levels', D2))])
table(s, ['Question', 'Measure'] + SHORT, rows, .89, 1.98, [1.6, 2.2, 1.9375, 1.9375, 1.9375, 1.9375], 4.0, size=14, head_size=14, head_h=.45, left_cols=(0, 1))
txt(s, 'Mean level error uses the probability-weighted expected level, not only the most likely level. The level ranges are decades, so one level is not a fixed physical error.',
    .95, 6.25, 11.4, .5, size=12.5, color=C['muted'])
W08 = {k: sum(level_counts('kev-0.8b', k, D2, 'truth_counts')[3:]) / D2M['kev-0.8b']['table'][k]['n'] for k in ['area', 'power', 'slack']}
assert all(math.isclose(W08[k], D2M['kev-0.8b']['table'][k]['within_one']) for k in W08)
notes(s, 'The saved scorer computes within_one from the most likely level against the true level: the two differ by at most one. The mean absolute level error (mae_levels) uses the probability-weighted expected level, sum(k × p_k) / sum(p_k), and averages |expected − true|; '
      'so it is not the error of the most likely level and need not follow the exact score. Counts: 92 area, 85 power, 92 frequency, 92 slack. '
      'KEV-0.8B with examples gives level 5 to every design for area, power, and slack, so its within-one score is the share of true levels 4 and 5: '
      + ', '.join(f'{k} {pct(W08[k])}' for k in W08) + '. All values come from the scorer and are verified against the raw answers with a tolerance of 1e-12.')

# A9. Pair orders.
s = slide('Pair questions: both orders, one order, neither', 'Counts per pair. Each pair has two answers, one per order. Each cell: no examples → three examples.', 'APPENDIX')
rows = []
for key, label in [('smaller_area', 'Area'), ('lower_power', 'Power'), ('higher_reach', 'Frequency')]:
    n = PAIRS[key]
    for stat, label2 in [('both_correct', 'both orders correct'), ('one_correct', 'one order correct'), ('neither_correct', 'neither order correct')]:
        rows.append([f'{label}: {label2}', n] + [change(PC[1][m][key][stat], PC[2][m][key][stat]) for m in M])
    assert all(sum(PC[e][m][key][st] for st in ['both_correct', 'one_correct', 'neither_correct']) == n for m in M for e in (1, 2))
table(s, ['Pair outcome', 'Pairs'] + SHORT, rows, .89, 1.98, [3.4, 1.0, 1.7875, 1.7875, 1.7875, 1.7875], 4.15, size=14, head_size=14, head_h=.45)
txt(s, 'The two correct letters of a pair are opposite, so “one order correct” means the same letter in both orders. A letter swap can still be wrong in both orders.',
    .95, 6.3, 11.4, .5, size=12.5, color=C['muted'])
notes(s, 'For a pair with opposite correct letters in the two orders, the two answers of a model fall into 0, 1, or 2 correct. The both-orders count is read from the paired saved answers, not estimated from percentages. '
      'One order correct equals the same letter in both orders (the scorer’s same_letter_both_orders rate times the number of pairs); it is not an accuracy measure. A consistent swap of the letter can still choose the wrong design in both orders. '
      'A model with one letter for every question has every pair in one order correct (KEV-0.8B area and frequency, KEV-4B frequency, with examples). The pairs share designs and problem families, so the questions are not independent samples. Each row triplet sums to the number of pairs (checked).')

# A10. Design groups.
GROUP_NAMES = {'metrex': 'MetRex designs', 'mutant': 'Injected synthesis errors (mutants)', 'natural': 'LLM-written candidates, natural result',
               'noncompile': 'Injected compile errors', 'reference': 'Human-written reference designs', 'revold': 'REvolution ASP-DAC 2026 candidates (pairs)',
               'revolution': 'ARES v0 designs', 'trap': 'Accepted edits and rejected twins (traps)'}
s = slide('Design groups in the test split', 'The retained designs by the group of the dataset plan. Pairs come from the REvolution candidates.', 'APPENDIX')
table(s, ['Design group', 'All retained', 'Test designs', 'Test synthesis failures'],
      [[GROUP_NAMES[g], v['all_designs'], v['test_designs'], v['test_synthesis_failures']] for g, v in GROUPS.items()],
      .89, 1.98, [5.55, 1.8, 1.9, 2.3], 3.85, size=15, head_size=14, head_h=.45)
txt(s, 'Not measured here: energy per operation, functional equivalence, post-route timing, measured chips, other model families, and the rename and comment controls of the earlier plan.',
    .95, 6.1, 11.4, .55, size=13, color=C['muted'])
notes(s, 'The group counts are based on the retained basic design records only, not on requests or pair questions. All retained designs total 151 and test designs 123. '
      'The plan ran 159 designs; 8 were excluded: 4 mutants with an unintended class, 2 trap twins that did not synthesize as required, and 2 noncompile designs whose DC class was not the intended one. '
      'Twelve retained performance designs across both splits have no usable power value, leaving 98 power questions overall and 85 in the test split. The dataset also skips pairs with incompatible testbenches or with a metric difference below 10%. '
      'Some groups contain near variants; the split per base design prevents a design and its edit from being in different splits (checked in analyse_kev.py). Names of the sources follow MV (2026-10-16): the REvolution v0 designs are ARES.',
      'data/questions/dataset_v1.card.md; notes/DATASET.md; harness/build.py; presentations/data/verified_metrics.json')

# A11. Condor runs.
GPU_NAME = {'NVIDIA GeForce RTX 5060 Ti': 'RTX 5060 Ti', 'NVIDIA GeForce RTX 3090': 'RTX 3090', 'NVIDIA RTX PRO 5000 Blackwell': 'RTX PRO 5000', 'NVIDIA A100-SXM4-80GB': 'A100 80 GB'}
MAN1 = [DM[m]['manifest'] for m in M]
MAN2 = [D2M[m]['manifest'] for m in M]
assert all(x['model_revision'] == y['model_revision'] for x, y in zip(MAN1, MAN2))
assert all(x['maximum_actual_input_tokens'] == D2['tokens']['max'] for x in MAN2)
s = slide('The Condor runs of the two experiments', 'Job times include the setup and differ in hardware: not a speed ranking. All requests answered, zero errors.', 'APPENDIX')
rows = [[LEG1, NAMES[i], x['job_id'], 'RTX 3080 Ti' if i == 0 else 'A100 80 GB', f"{x['condor_execution_seconds']:.0f} s", f"{x['median_model_latency_ms']:.0f} ms"] for i, x in enumerate(MAN1)]
rows += [[LEG2, NAMES[i], x['job_id'], GPU_NAME[x['gpu']], f"{x['condor_execution_seconds']:.0f} s", f"{x['median_model_latency_ms']:.0f} ms"] for i, x in enumerate(MAN2)]
table(s, ['Experiment', 'Model', 'Job', 'GPU', 'Job time', 'Median model time per request'], rows, .89, 1.98, [1.9, 1.5, .9, 2.0, 1.6, 3.65], 4.1, size=14, head_size=14, head_h=.45, left_cols=(0, 1))
txt(s, f'No examples: 348 requests per model (both splits), maximum input {max(x["maximum_actual_input_tokens"] for x in MAN1):,} tokens. Three examples: 725 requests per model, maximum input {D2["tokens"]["max"]:,} tokens; {len(D2["tokens"]["over_8192"])} requests exceed 8,192 tokens, the validated context of 0.8B, 4B, and 9B.',
    .95, 6.2, 11.4, .6, size=12.5, color=C['muted'])
JOBS1, JOBS2 = ', '.join(x['job_id'] for x in MAN1), ', '.join(x['job_id'] for x in MAN2)
EXEC1, EXEC2 = [', '.join(f"{x['condor_execution_seconds']:.0f}" for x in man) for man in (MAN1, MAN2)]
LAT1, LAT2 = [', '.join(f"{x['median_model_latency_ms']:.1f}" for x in man) for man in (MAN1, MAN2)]
HOSTS2 = ', '.join(x['host'].split('.')[0] for x in MAN2)
MEM2 = ', '.join(f"{x['maximum_gpu_memory_of_server_mib']:,}" for x in MAN2)
OVER = ', '.join(q.split(':', 2)[2] for q in D2['tokens']['over_8192'])
notes(s, f'Experiment 1: Condor jobs {JOBS1} completed with return code 0 on 15 October 2026; job execution seconds {EXEC1} (setup, smoke file, and inference, without the queue wait); '
      f'median server-reported model times {LAT1} ms over the 348 requests. KEV-0.8B ran on euler (RTX 3080 Ti); the larger sizes on aristoteles (A100-SXM4-80GB). '
      f'Experiment 2: jobs {JOBS2} on {HOSTS2}; execution seconds {EXEC2}; '
      f'median model times {LAT2} ms over the 725 requests; peak GPU memory of the Kev server {MEM2} MiB. '
      f'Every model used the same v1.0 checkpoint revision in both experiments. The requests above 8,192 tokens are {OVER}; Kev-27B is validated to 65,536 tokens. No dollar cost is measured.',
      'results/kev/*/manifest.json; condor/logs/*.log')

# A12. Files and checksums.
s = slide('Files and checksums for reproduction', 'Every number is tied to the frozen question files and the saved answers. No model was rerun for this deck.', 'APPENDIX')
table(s, ['Model', 'Job, no examples', 'Completed', 'Job, three examples', 'Completed', 'Checkpoint revision (prefix)'],
      [[NAMES[i], x['job_id'], x['completion_time'][5:16], y['job_id'], y['completion_time'][5:16], x['model_revision'][:12]] for i, (x, y) in enumerate(zip(MAN1, MAN2))],
      .89, 1.98, [1.5, 1.6, 1.9, 1.75, 1.9, 2.9], 2.4, size=14, head_size=13, head_h=.45)
txt(s, 'Inputs and evidence', 1.05, 4.65, 11.1, .4, size=20, bold=True)
txt(s, 'Questions: data/questions/dataset_v1.jsonl and dataset_v1_fewshot3.jsonl (harness/fewshot.py)\n'
       'Answers, scores, manifests, and comparisons: results/kev/dataset_v1_<model>_2026101[45]/ and results/kev/dataset_v1_fewshot3_<model>_20261016/\n'
       'Verified tables and the build: presentations/data/verified_metrics.json, verified_metrics_fewshot3.json, compare_fewshot3.csv; presentations/src/analyse_kev.py, build_kev_findings.py',
    1.05, 5.15, 11.1, 1.0, size=14)
notes(s, 'Full provenance:\n' + json.dumps({m: {'job_1': DM[m]['manifest']['job_id'], 'job_2': D2M[m]['manifest']['job_id'], 'revision': DM[m]['manifest']['model_revision'],
                                               'answers_1': DM[m]['answers_file'], 'answers_1_sha256': DM[m]['answers_sha256'],
                                               'answers_2': D2M[m]['answers_file'], 'answers_2_sha256': D2M[m]['answers_sha256']} for m in M}, indent=2) +
      '\nAll times are the local time of mercury (CEST) as recorded in the manifests. The build verifies unique request IDs, HTTP 200 for every answer, the answer key of each question, no shared base design between the splits, the dataset fingerprints, '
      'the server input tokens against the builder counts (experiment 2), the smoke answers, the saved scores against freshly recomputed scores (tolerance 1e-12), and every paired comparison against compare_with_dataset_v1.json. '
      'The theme and the slide masters come from the supplied TinyTapeOut reference deck.')


# ----------------------------------------------------------------------------------------------
# Validate and save
# ----------------------------------------------------------------------------------------------
def validate_and_save():
    assert not WARN, '\n'.join(WARN)
    issues = []
    for i, s in enumerate(prs.slides, 1):
        assert s.has_notes_slide and len(s.notes_slide.notes_text_frame.text) > 200
        for sh in s.shapes:
            if sh.left < -100 or sh.top < -100 or sh.left + sh.width > prs.slide_width + 100 or sh.top + sh.height > prs.slide_height + 100:
                issues.append(f'Slide {i}: out-of-bounds shape {sh.name}')
    assert not issues, issues
    DEST.mkdir(exist_ok=True)
    if OUTPUT.exists():
        bak = DEST / 'backups'
        bak.mkdir(exist_ok=True)
        shutil.copy2(OUTPUT, bak / (OUTPUT.stem + '_' + datetime.now().strftime('%Y%m%d_%H%M%S') + '.pptx'))
    prs.save(OUTPUT)
    check = Presentation(OUTPUT)
    assert len(check.slides) == 29 == len(SLIDE_META)
    assert all(s.has_notes_slide for s in check.slides)
    (DEST / 'qa/slide_manifest.json').write_text(json.dumps(SLIDE_META, indent=2) + '\n')
    (DEST / 'qa/text_boxes.json').write_text(json.dumps(TEXT_BOXES, indent=2) + '\n')
    print(f'Saved {OUTPUT.relative_to(ROOT)}: {len(check.slides)} slides; all shapes in bounds; notes on every slide.')


if __name__ == '__main__':
    validate_and_save()
